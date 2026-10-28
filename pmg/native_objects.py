"""Compile shared visual objects into native editable PowerPoint content.

Content construction only: all animation XML stays in pmg.animations, Morph
stays in shared transition helpers, and delivery stays with PMGOrchestrator.
"""
from typing import Dict, Optional
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from pmg.layers import LayerTree

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def _custom_geometry(shape, paths):
    properties = shape._element.spPr
    for node in list(properties):
        if node.tag in (f'{{{A_NS}}}prstGeom', f'{{{A_NS}}}custGeom'):
            properties.remove(node)
    geometry = etree.Element(f'{{{A_NS}}}custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        etree.SubElement(geometry, f'{{{A_NS}}}{tag}')
    etree.SubElement(geometry, f'{{{A_NS}}}rect', l='0', t='0', r='r', b='b')
    path_list = etree.SubElement(geometry, f'{{{A_NS}}}pathLst')
    path_node = etree.SubElement(path_list, f'{{{A_NS}}}path', w='100000', h='100000')
    for points in paths:
        for index, (x, y) in enumerate(points):
            command = etree.SubElement(path_node, f'{{{A_NS}}}{"moveTo" if index == 0 else "lnTo"}')
            etree.SubElement(command, f'{{{A_NS}}}pt', x=str(round(x*100000)), y=str(round(y*100000)))
        etree.SubElement(path_node, f'{{{A_NS}}}close')
    # Geometry precedes fill/line/effects in CT_ShapeProperties.
    properties.insert(1 if properties.find(f'{{{A_NS}}}xfrm') is not None else 0, geometry)


def fit_layer_bounds(tree: LayerTree, bounds):
    left, top, width, height = bounds
    source_w, source_h = tree.source_size_px
    ratio = min(width/source_w, height/source_h)
    return left+(width-source_w*ratio)/2, top+(height-source_h*ratio)/2, ratio


def compile_layer_tree(slide, tree: LayerTree, bounds=(.8, 1.6, 11.7, 5.0), include_ids=None):
    audit = tree.audit()
    if not audit['is_valid']:
        raise ValueError('Invalid layer tree: '+str(audit['defects']))
    if tree.reassembly and tree.reassembly.get('status') != 'PASS':
        raise ValueError('Layer reassembly must pass before native compilation')
    origin_x, origin_y, factor = fit_layer_bounds(tree, bounds)
    objects = {obj.object_id: obj for obj in tree.objects}
    included = set(include_ids) if include_ids is not None else set(objects)
    mapping: Dict[str, object] = {}

    def needed(obj):
        return obj.object_id in included or any(needed(objects[child]) for child in obj.children)

    def compile_object(obj, collection):
        if obj.kind == 'group':
            children = [objects[child] for child in obj.children if needed(objects[child])]
            if not children:
                return
            shape = collection.add_group_shape()
            shape.name = '!!'+obj.object_id
            mapping[obj.object_id] = shape
            for child in sorted(children, key=lambda item: item.z_order):
                compile_object(child, shape.shapes)
            return
        scale = factor if obj.units == 'px' else 1.0
        x, y, width, height = obj.geometry
        geometry = [Inches(origin_x+x*scale), Inches(origin_y+y*scale), Inches(width*scale), Inches(height*scale)]
        if obj.kind == 'image':
            shape = collection.add_picture(obj.media_path, *geometry)
            shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom = obj.crop
        elif obj.kind == 'text':
            shape = collection.add_textbox(*geometry)
            shape.text_frame.text = obj.text
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.size = Pt(obj.metadata.get('font_size_pt', 18))
                paragraph.font.color.rgb = RGBColor.from_string(obj.fill)
        elif obj.kind in ('rect', 'ellipse', 'polygon'):
            shape = collection.add_shape(MSO_SHAPE.OVAL if obj.kind == 'ellipse' else MSO_SHAPE.RECTANGLE, *geometry)
            if obj.kind == 'polygon':
                _custom_geometry(shape, obj.paths)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(obj.fill.lstrip('#'))
            if obj.opacity < 1:
                color = shape._element.spPr.find(f'{{{A_NS}}}solidFill/{{{A_NS}}}srgbClr')
                etree.SubElement(color, f'{{{A_NS}}}alpha', val=str(round(obj.opacity*100000)))
            if obj.stroke:
                shape.line.color.rgb = RGBColor.from_string(obj.stroke.lstrip('#'))
                shape.line.width = Pt(obj.stroke_width)
            else:
                shape.line.fill.background()
        else:
            raise ValueError(f'Native layer compiler does not support {obj.kind!r}; use a compatible media/model adapter')
        shape.name = '!!'+obj.object_id
        shape.rotation = obj.rotation_deg
        mapping[obj.object_id] = shape

    compile_object(objects[tree.root_id], slide.shapes)
    return mapping
