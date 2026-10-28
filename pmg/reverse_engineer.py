"""
pmg.reverse_engineer
====================
PPTX Motion Reverse Engineering & Animation Style Transfer Engine (v1.4.0).

Analyzes an existing animated PowerPoint presentation (single slide or slide range),
extracts its complete object hierarchy, geometry, z-order, text/image properties,
group structure, transitions, Morph anchors, and native <p:timing> animation tree into
a reusable PresentationMotionSpec, and reconstructs or transfers those motion principles
onto new editable PowerPoint slides with full Animation Pane compatibility.
"""

import os
import json
import math
from dataclasses import dataclass, field, asdict
from typing import List, Dict, Any, Optional, Tuple

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

from pmg.animations import (
    SlideAnimationTimeline,
    P_NS,
    A_NS,
    NSMAP,
    _xpath,
    verify_animation_pane_com,
)
from pmg.validator import parse_slide_transition, EMU_PER_INCH
from pmg.engine import MotionDeck, Themes


PRST_TO_MSO = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "star4": MSO_SHAPE.STAR_4_POINT,
    "diamond": MSO_SHAPE.DIAMOND,
    "chevron": MSO_SHAPE.CHEVRON,
    "hexagon": MSO_SHAPE.HEXAGON,
}


@dataclass
class ExtractedObjectSpec:
    spid: int
    name: str
    role: str
    object_type: str
    prst_geom: str
    z_order: int
    geometry_in: Tuple[float, float, float, float]  # (x, y, w, h)
    rotation_deg: float = 0.0
    opacity: float = 1.0
    fill_hex: str = "FFFFFF"
    line_hex: str = "CBD5E1"
    has_fill: bool = True
    has_line: bool = True
    text_content: str = ""
    font_size_pt: float = 14.0
    text_color_hex: str = "0F172A"
    is_morph_anchor: bool = False
    children: List[Dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class ExtractedAnimationSpec:
    order_index: int
    parallel_group_id: int
    target_spid: int
    target_name: str
    target_role: str
    effect_type: str
    preset_class: str
    preset_id: int
    preset_subtype: int
    direction: str
    trigger: str
    duration_ms: int
    delay_ms: int
    motion_path: Optional[str] = None
    rotation_deg: float = 0.0
    scale_factor: Tuple[float, float] = (1.0, 1.0)
    opacity_to: float = 1.0
    repeat_count: float = 1.0
    auto_reverse: bool = False

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class SlideBlueprintSpec:
    slide_number: int
    transition_type: str
    transition_duration_ms: int
    objects: List[ExtractedObjectSpec] = field(default_factory=list)
    animations: List[ExtractedAnimationSpec] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "slide_number": self.slide_number,
            "transition_type": self.transition_type,
            "transition_duration_ms": self.transition_duration_ms,
            "objects": [o.to_dict() for o in self.objects],
            "animations": [a.to_dict() for a in self.animations],
        }


@dataclass
class PresentationMotionSpec:
    source_pptx: str
    slide_width_in: float
    slide_height_in: float
    slides: List[SlideBlueprintSpec] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "source_pptx": self.source_pptx,
            "slide_width_in": self.slide_width_in,
            "slide_height_in": self.slide_height_in,
            "slides": [s.to_dict() for s in self.slides],
        }

    def save(self, path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(self.to_dict(), f, indent=2, ensure_ascii=False)
        return path


def _infer_role(name: str, obj_type: str, geom: Tuple[float, float, float, float], prst: str) -> str:
    nl = name.lower()
    x, y, w, h = geom
    if "badge" in nl or (h < 0.65 and y < 1.2):
        return "badge"
    if "title" in nl or (obj_type == "textbox" and y < 2.2):
        return "title"
    if "phone" in nl or "device" in nl or "hero" in nl or "globe" in nl:
        return "hero_device"
    if "status" in nl or "telemetry" in nl or (y > 5.8 and w > 6.0):
        return "status_bar"
    if "packet" in nl or "orb" in nl or "node" in nl or prst in ("ellipse", "star4"):
        return "accent_node"
    if w > 11.0 and h > 3.0:
        return "backdrop"
    if "card" in nl or "widget" in nl or (w >= 1.5 and h >= 1.0):
        return "ui_card"
    return "element"


def _hex_to_rgb(hex_str: str) -> RGBColor:
    h = (hex_str or "FFFFFF").lstrip("#")
    if len(h) != 6:
        h = "FFFFFF"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class PPTXMotionReverseEngineer:
    """
    Reverse-engineers object structure, geometry, z-order, transitions, and native
    <p:timing> animations from existing PowerPoint slides and transfers them onto
    new editable presentations.
    """

    def extract_motion_spec(
        self,
        pptx_path: str,
        slide_range: Optional[Tuple[int, int]] = None,
    ) -> PresentationMotionSpec:
        prs = pptx.Presentation(pptx_path)
        w_in = round(prs.slide_width / EMU_PER_INCH, 3)
        h_in = round(prs.slide_height / EMU_PER_INCH, 3)

        total = len(prs.slides)
        start_s, end_s = slide_range if slide_range else (1, total)
        start_s = max(1, min(total, start_s))
        end_s = max(start_s, min(total, end_s))

        slide_specs: List[SlideBlueprintSpec] = []

        for s_num in range(start_s, end_s + 1):
            slide = prs.slides[s_num - 1]
            xml_text = etree.tostring(slide._element, encoding="unicode")
            trans_type, trans_dur, _ = parse_slide_transition(xml_text, s_num)

            extracted_objs: List[ExtractedObjectSpec] = []
            obj_by_spid: Dict[int, ExtractedObjectSpec] = {}

            for z_idx, shp in enumerate(slide.shapes):
                obj_spec = self._extract_shape_spec(shp, z_idx)
                extracted_objs.append(obj_spec)
                obj_by_spid[obj_spec.spid] = obj_spec

            extracted_anims = self._extract_slide_animations(slide, obj_by_spid)

            slide_specs.append(
                SlideBlueprintSpec(
                    slide_number=s_num,
                    transition_type=trans_type,
                    transition_duration_ms=trans_dur,
                    objects=extracted_objs,
                    animations=extracted_anims,
                )
            )

        return PresentationMotionSpec(
            source_pptx=os.path.abspath(pptx_path),
            slide_width_in=w_in,
            slide_height_in=h_in,
            slides=slide_specs,
        )

    def _extract_shape_spec(self, shp, z_order: int) -> ExtractedObjectSpec:
        spid = shp.shape_id
        name = shp.name or f"Shape_{spid}"
        el = shp._element

        x = round(shp.left / EMU_PER_INCH, 4) if getattr(shp, "left", None) is not None else 0.0
        y = round(shp.top / EMU_PER_INCH, 4) if getattr(shp, "top", None) is not None else 0.0
        w = round(shp.width / EMU_PER_INCH, 4) if getattr(shp, "width", None) is not None else 1.0
        h = round(shp.height / EMU_PER_INCH, 4) if getattr(shp, "height", None) is not None else 1.0

        rot_deg = 0.0
        xfrms = _xpath(el, ".//a:xfrm[@rot]")
        if xfrms:
            try:
                rot_deg = round(int(xfrms[0].get("rot", "0")) / 60000.0, 2)
            except ValueError:
                pass

        prst_geom = "rect"
        prsts = _xpath(el, ".//a:prstGeom[@prst]")
        if prsts:
            prst_geom = prsts[0].get("prst", "rect")

        # Determine object type
        tag = el.tag.split("}")[-1] if "}" in el.tag else el.tag
        if tag == "grpSp":
            obj_type = "group"
        elif tag == "pic":
            obj_type = "picture"
        elif tag == "cxnSp":
            obj_type = "connector"
        else:
            txBody = _xpath(el, "./p:txBody")
            noFill = _xpath(el, "./p:spPr/a:noFill")
            if txBody and (noFill or "TextBox" in name):
                obj_type = "textbox"
            else:
                obj_type = "shape"

        # Fill & alpha
        has_fill = not bool(_xpath(el, "./p:spPr/a:noFill"))
        fill_hex = "F8FAFC"
        opacity = 1.0
        srgb_fills = _xpath(el, "./p:spPr/a:solidFill/a:srgbClr")
        if srgb_fills:
            fill_hex = srgb_fills[0].get("val", "F8FAFC")
            alphas = _xpath(srgb_fills[0], "./a:alpha")
            if alphas:
                try:
                    opacity = round(int(alphas[0].get("val", "100000")) / 100000.0, 3)
                except ValueError:
                    pass

        # Line
        has_line = not bool(_xpath(el, "./p:spPr/a:ln/a:noFill"))
        line_hex = "CBD5E1"
        srgb_lines = _xpath(el, "./p:spPr/a:ln/a:solidFill/a:srgbClr")
        if srgb_lines:
            line_hex = srgb_lines[0].get("val", "CBD5E1")

        # Text
        text_content = ""
        font_size_pt = 14.0
        text_color_hex = "0F172A"
        if getattr(shp, "has_text_frame", False):
            text_content = shp.text_frame.text.strip()
            rPrs = _xpath(el, ".//a:rPr[@sz]")
            if rPrs:
                try:
                    font_size_pt = round(int(rPrs[0].get("sz", "1400")) / 100.0, 1)
                except ValueError:
                    pass
            t_clrs = _xpath(el, ".//a:rPr/a:solidFill/a:srgbClr")
            if t_clrs:
                text_color_hex = t_clrs[0].get("val", "0F172A")

        # Group children
        children = []
        if obj_type == "group" and hasattr(shp, "shapes"):
            for c_idx, child_shp in enumerate(shp.shapes):
                children.append(self._extract_shape_spec(child_shp, c_idx).to_dict())

        role = _infer_role(name, obj_type, (x, y, w, h), prst_geom)

        return ExtractedObjectSpec(
            spid=spid,
            name=name,
            role=role,
            object_type=obj_type,
            prst_geom=prst_geom,
            z_order=z_order,
            geometry_in=(x, y, w, h),
            rotation_deg=rot_deg,
            opacity=opacity,
            fill_hex=fill_hex,
            line_hex=line_hex,
            has_fill=has_fill,
            has_line=has_line,
            text_content=text_content,
            font_size_pt=font_size_pt,
            text_color_hex=text_color_hex,
            is_morph_anchor=name.startswith("!!"),
            children=children,
        )

    def _extract_slide_animations(
        self,
        slide,
        obj_by_spid: Dict[int, ExtractedObjectSpec],
    ) -> List[ExtractedAnimationSpec]:
        timing = slide._element.find(f"{{{P_NS}}}timing")
        if timing is None:
            return []

        effect_nodes = _xpath(
            timing,
            './/p:cTn[@presetClass="entr" or @presetClass="emph" or @presetClass="path" or @presetClass="exit"]',
        )

        anims: List[ExtractedAnimationSpec] = []
        parallel_group_id = 0

        for idx, eff_cTn in enumerate(effect_nodes):
            p_class = eff_cTn.get("presetClass", "entr")
            p_id = int(eff_cTn.get("presetID", "10"))
            p_sub = int(eff_cTn.get("presetSubtype", "0"))
            node_type = eff_cTn.get("nodeType", "withEffect")
            auto_rev = eff_cTn.get("autoRev") == "1"
            rep_raw = eff_cTn.get("repeatCount")
            rep_count = (int(rep_raw) / 1000.0) if rep_raw and rep_raw.isdigit() else 1.0

            if idx == 0 or node_type in ("afterEffect", "clickEffect"):
                parallel_group_id += 1

            delay_ms = 0
            conds = _xpath(eff_cTn, "./p:stCondLst/p:cond[@delay]")
            if conds:
                d_val = conds[0].get("delay", "0")
                if d_val.isdigit():
                    delay_ms = int(d_val)

            sp_tgts = _xpath(eff_cTn, ".//p:spTgt[@spid]")
            if not sp_tgts:
                continue
            spid = int(sp_tgts[0].get("spid"))
            obj_spec = obj_by_spid.get(spid)
            target_name = obj_spec.name if obj_spec else f"Shape_{spid}"
            target_role = obj_spec.role if obj_spec else "element"

            dur_ms = 500
            for c in _xpath(eff_cTn, ".//p:cBhvr/p:cTn[@dur]"):
                d = c.get("dur", "500")
                if d.isdigit() and int(d) > 1:
                    dur_ms = int(d)
                    break

            sub_dir_map = {4: "up", 1: "down", 2: "left", 8: "right"}
            direction = sub_dir_map.get(p_sub, "up")

            effect_type = self._map_preset_to_effect_type(p_class, p_id)

            motion_path = None
            anim_motions = _xpath(eff_cTn, ".//p:animMotion[@path]")
            if anim_motions:
                motion_path = anim_motions[0].get("path")
                effect_type = "motion_path"

            rotation_deg = 0.0
            anim_rots = _xpath(eff_cTn, ".//p:animRot[@by]")
            if anim_rots:
                by_v = anim_rots[0].get("by", "0")
                if by_v.lstrip("-").isdigit():
                    rotation_deg = round(int(by_v) / 60000.0, 2)
                effect_type = "rotate"

            scale_factor = (1.0, 1.0)
            anim_scales = _xpath(eff_cTn, ".//p:animScale/p:by")
            if anim_scales:
                bx = int(anim_scales[0].get("x", "100000")) / 100000.0
                by = int(anim_scales[0].get("y", "100000")) / 100000.0
                scale_factor = (round(bx, 3), round(by, 3))
                effect_type = "scale"

            trig_map = {"withEffect": "with_previous", "afterEffect": "after_previous", "clickEffect": "on_click"}
            trigger = trig_map.get(node_type, "with_previous")

            anims.append(
                ExtractedAnimationSpec(
                    order_index=idx + 1,
                    parallel_group_id=parallel_group_id,
                    target_spid=spid,
                    target_name=target_name,
                    target_role=target_role,
                    effect_type=effect_type,
                    preset_class=p_class,
                    preset_id=p_id,
                    preset_subtype=p_sub,
                    direction=direction,
                    trigger=trigger,
                    duration_ms=dur_ms,
                    delay_ms=delay_ms,
                    motion_path=motion_path,
                    rotation_deg=rotation_deg,
                    scale_factor=scale_factor,
                    repeat_count=rep_count,
                    auto_reverse=auto_rev,
                )
            )

        return anims

    @staticmethod
    def _map_preset_to_effect_type(preset_class: str, preset_id: int) -> str:
        if preset_class == "entr":
            return {1: "appear", 2: "fly_in", 10: "fade", 22: "wipe", 23: "zoom"}.get(preset_id, "fade")
        if preset_class == "emph":
            return {6: "scale", 8: "rotate", 9: "transparency"}.get(preset_id, "scale")
        if preset_class == "path":
            return "motion_path"
        if preset_class == "exit":
            return {1: "disappear", 2: "fly_out", 10: "fade_out", 22: "wipe_out"}.get(preset_id, "fade_out")
        return "fade"

    def reconstruct_and_transfer(
        self,
        motion_spec: PresentationMotionSpec,
        output_pptx_path: str,
        content_transfer: Optional[Dict[str, str]] = None,
        palette_transfer: Optional[Dict[str, str]] = None,
    ) -> str:
        """
        Reconstructs an editable PowerPoint presentation from a PresentationMotionSpec,
        optionally transferring the extracted animation principles onto new content and colors.
        """
        content_transfer = content_transfer or {}
        palette_transfer = palette_transfer or {}

        deck = MotionDeck(theme=Themes.LIGHT_ACADEMIC_CYAN)

        for s_spec in motion_spec.slides:
            slide = deck.add_slide(
                transition=s_spec.transition_type if s_spec.slide_number > 1 else "none",
                duration_ms=s_spec.transition_duration_ms or 1600,
            )
            tl = SlideAnimationTimeline(slide_number=s_spec.slide_number)
            new_shape_by_old_spid: Dict[int, Any] = {}

            for obj in sorted(s_spec.objects, key=lambda o: o.z_order):
                x, y, w, h = obj.geometry_in
                w = max(0.15, w)
                h = max(0.15, h)

                if obj.object_type == "textbox":
                    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                else:
                    mso_type = PRST_TO_MSO.get(obj.prst_geom, MSO_SHAPE.ROUNDED_RECTANGLE)
                    shp = slide.shapes.add_shape(mso_type, Inches(x), Inches(y), Inches(w), Inches(h))

                shp.name = content_transfer.get(f"name:{obj.name}", obj.name)

                # Apply fill & line
                if obj.object_type != "textbox":
                    if obj.has_fill:
                        shp.fill.solid()
                        hex_col = palette_transfer.get(obj.role, obj.fill_hex)
                        shp.fill.fore_color.rgb = _hex_to_rgb(hex_col)
                    else:
                        shp.fill.background()

                    if obj.has_line:
                        shp.line.color.rgb = _hex_to_rgb(obj.line_hex)
                        shp.line.width = Pt(1.25)
                    else:
                        shp.line.fill.background()

                # Apply text (with optional content transfer)
                new_text = content_transfer.get(obj.name, content_transfer.get(obj.role, obj.text_content))
                if new_text and getattr(shp, "has_text_frame", False):
                    tf = shp.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    p.text = new_text
                    p.font.size = Pt(obj.font_size_pt)
                    p.font.bold = obj.role in ("badge", "title", "hero_device")
                    p.font.color.rgb = _hex_to_rgb(obj.text_color_hex)

                # Apply rotation
                if obj.rotation_deg != 0.0:
                    try:
                        shp._element.spPr.xfrm.set("rot", str(int(obj.rotation_deg * 60000)))
                    except Exception:
                        pass

                new_shape_by_old_spid[obj.spid] = shp

            # Reconstruct the complete native animation timeline in exact order & parallel grouping
            for anim in s_spec.animations:
                target_shp = new_shape_by_old_spid.get(anim.target_spid)
                if target_shp is None:
                    continue
                tl.add_effect(
                    shape=target_shp,
                    effect_type=anim.effect_type,
                    trigger=anim.trigger,
                    duration_ms=anim.duration_ms,
                    delay_ms=anim.delay_ms,
                    direction=anim.direction,
                    motion_path=anim.motion_path,
                    rotation_deg=anim.rotation_deg,
                    scale_factor=anim.scale_factor,
                    opacity_to=anim.opacity_to,
                    repeat_count=anim.repeat_count,
                    auto_reverse=anim.auto_reverse,
                )

            if tl.steps:
                tl.compile_to_openxml(slide)

        deck.save(output_pptx_path)
        return os.path.abspath(output_pptx_path)

    def compare_reference_and_reconstructed(
        self,
        ref_pptx: str,
        recon_pptx: str,
    ) -> Dict[str, Any]:
        """
        Compares the extracted motion specifications of the reference and reconstructed
        presentations, identifying exact matches and approximations across objects and animations.
        """
        ref_spec = self.extract_motion_spec(ref_pptx)
        rec_spec = self.extract_motion_spec(recon_pptx)

        exact_matches: List[Dict[str, Any]] = []
        approximations: List[Dict[str, Any]] = []
        position_errors_in: List[float] = []
        duration_errors_ms: List[int] = []

        for s_ref, s_rec in zip(ref_spec.slides, rec_spec.slides):
            for o_ref, o_rec in zip(s_ref.objects, s_rec.objects):
                dx = o_rec.geometry_in[0] - o_ref.geometry_in[0]
                dy = o_rec.geometry_in[1] - o_ref.geometry_in[1]
                pos_err = math.hypot(dx, dy)
                position_errors_in.append(pos_err)

            for a_ref, a_rec in zip(s_ref.animations, s_rec.animations):
                dur_err = abs(a_rec.duration_ms - a_ref.duration_ms)
                duration_errors_ms.append(dur_err)
                if (
                    a_ref.effect_type == a_rec.effect_type
                    and a_ref.trigger == a_rec.trigger
                    and dur_err == 0
                    and a_ref.motion_path == a_rec.motion_path
                ):
                    exact_matches.append({
                        "order_index": a_ref.order_index,
                        "effect_type": a_ref.effect_type,
                        "trigger": a_ref.trigger,
                        "duration_ms": a_ref.duration_ms,
                    })
                else:
                    approximations.append({
                        "order_index": a_ref.order_index,
                        "ref_effect": a_ref.effect_type,
                        "recon_effect": a_rec.effect_type,
                        "duration_diff_ms": dur_err,
                    })

        return {
            "reference_slides": len(ref_spec.slides),
            "reconstructed_slides": len(rec_spec.slides),
            "reference_object_count": sum(len(s.objects) for s in ref_spec.slides),
            "reconstructed_object_count": sum(len(s.objects) for s in rec_spec.slides),
            "reference_animation_count": sum(len(s.animations) for s in ref_spec.slides),
            "reconstructed_animation_count": sum(len(s.animations) for s in rec_spec.slides),
            "exact_animation_matches": len(exact_matches),
            "approximated_animations": len(approximations),
            "mean_position_error_in": round(sum(position_errors_in) / max(1, len(position_errors_in)), 4),
            "mean_duration_error_ms": round(sum(duration_errors_ms) / max(1, len(duration_errors_ms)), 2),
            "exact_matches_detail": exact_matches,
            "approximations_detail": approximations,
        }


def build_v1_4_reference_smartphone_slide(output_path: str) -> str:
    """
    Creates the canonical v1.4 reference animated PPTX slide:
    A smartphone device frame flying into view, rotating, scaling up, and revealing
    interface components, data packet trajectories, and telemetry status.
    """
    deck = MotionDeck(theme=Themes.LIGHT_MODERN_STUDIO)
    slide = deck.add_slide(transition="none")
    tl = SlideAnimationTimeline(slide_number=1)

    # 1. Backdrop
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.name = "Studio_Backdrop"
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bg.line.fill.background()

    # 2. Badge
    badge = deck.add_badge(
        slide,
        "REFERENCE SLIDE · SMARTPHONE UI LAUNCH",
        left=0.85,
        top=0.6,
        width=4.2,
        height=0.38,
        bg_color=RGBColor(2, 132, 199),
        text_color=RGBColor(255, 255, 255),
        name="Ref_Badge",
    )

    # 3. Title
    title = deck.add_hero_title(
        slide,
        "EcoPhone Pro · Mobile Grid Controller",
        subtitle="Reference slide with smartphone fly-in, rotation, scale-up, and progressive UI reveal.",
        left=0.85,
        top=1.15,
        width=6.8,
        height=1.35,
        title_size=24,
        subtitle_size=11.5,
        title_color=RGBColor(15, 23, 42),
        name="Ref_Title",
    )

    # 4. Smartphone Outer Frame (Hero Device)
    phone_frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(2.65), Inches(3.2), Inches(4.1))
    phone_frame.name = "Hero_Smartphone_Frame"
    phone_frame.fill.solid()
    phone_frame.fill.fore_color.rgb = RGBColor(15, 23, 42)
    phone_frame.line.color.rgb = RGBColor(56, 189, 248)
    phone_frame.line.width = Pt(2.5)
    phone_frame.text_frame.text = "ECOPHONE OS v4"
    phone_frame.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    phone_frame.text_frame.paragraphs[0].font.size = Pt(13)

    # 5. Smartphone Screen Display
    phone_screen = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(3.15), Inches(2.8), Inches(3.3))
    phone_screen.name = "Smartphone_Display_Screen"
    phone_screen.fill.solid()
    phone_screen.fill.fore_color.rgb = RGBColor(224, 242, 254)
    phone_screen.line.color.rgb = RGBColor(2, 132, 199)
    phone_screen.text_frame.text = "LIVE SOLAR SYNC: 99.2%"
    phone_screen.text_frame.paragraphs[0].font.color.rgb = RGBColor(3, 105, 161)
    phone_screen.text_frame.paragraphs[0].font.size = Pt(12)

    # 6. UI Component Card 1
    ui_card_1, _ = deck.add_metric_card(
        slide,
        title="BATTERY HEALTH",
        value="99.8%",
        desc="Solid-state cell telemetry.",
        left=5.0,
        top=2.75,
        width=3.6,
        height=1.75,
        accent_color=RGBColor(5, 150, 105),
        name="UI_Component_Card_1",
    )

    # 7. UI Component Card 2
    ui_card_2, _ = deck.add_metric_card(
        slide,
        title="NEURAL MODEM",
        value="4.2 GBPS",
        desc="Ultra-low-power satellite link.",
        left=8.9,
        top=2.75,
        width=3.6,
        height=1.75,
        accent_color=RGBColor(2, 132, 199),
        name="UI_Component_Card_2",
    )

    # 8. UI Component Card 3
    ui_card_3, _ = deck.add_metric_card(
        slide,
        title="CARBON SAVED",
        value="-1.4 TONS",
        desc="Smart home load shifting.",
        left=5.0,
        top=4.8,
        width=7.5,
        height=1.4,
        accent_color=RGBColor(16, 185, 129),
        name="UI_Component_Card_3",
    )

    # 9. Sync Signal Node (travels along curved path from phone to UI cards)
    signal_orb = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.35), Inches(3.35), Inches(0.5), Inches(0.5))
    signal_orb.name = "Signal_Sync_Orb"
    signal_orb.fill.solid()
    signal_orb.fill.fore_color.rgb = RGBColor(14, 165, 233)
    signal_orb.line.color.rgb = RGBColor(255, 255, 255)

    # 10. Footer Status Bar
    footer = deck.add_badge(
        slide,
        "DEVICE TELEMETRY: ALL INTERFACE COMPONENTS ACTIVE",
        left=5.0,
        top=6.4,
        width=7.5,
        height=0.42,
        bg_color=RGBColor(15, 23, 42),
        text_color=RGBColor(255, 255, 255),
        name="Footer_Telemetry_Bar",
    )

    # Animation Timeline (15 native PowerPoint effects)
    tl.add_effect(badge, "appear", trigger="with_previous", duration_ms=150)
    tl.add_effect(title, "fade", trigger="with_previous", duration_ms=300, delay_ms=50)

    # Smartphone flies into view, rotates, and scales up!
    tl.add_effect(phone_frame, "fly_in", trigger="after_previous", direction="left", duration_ms=400)
    tl.add_effect(phone_screen, "fade", trigger="with_previous", duration_ms=350, delay_ms=50)
    tl.add_effect(phone_frame, "rotate", trigger="after_previous", rotation_deg=360.0, duration_ms=450)
    tl.add_effect(phone_frame, "scale", trigger="with_previous", scale_factor=(1.08, 1.08), duration_ms=350, auto_reverse=True)
    tl.add_effect(phone_screen, "scale", trigger="with_previous", scale_factor=(1.08, 1.08), duration_ms=350, auto_reverse=True)

    # Signal orb appears and travels along curved Bezier path to UI cards
    tl.add_effect(signal_orb, "appear", trigger="after_previous", duration_ms=100)
    tl.add_effect(signal_orb, "motion_path", trigger="with_previous", motion_path="M 0 0 C 0.08 -0.10 0.16 0.10 0.26 0 E", duration_ms=500)

    # Reveal interface components sequentially & simultaneously
    tl.add_effect(ui_card_1, "fly_in", trigger="after_previous", direction="up", duration_ms=350)
    tl.add_effect(ui_card_2, "fly_in", trigger="with_previous", direction="up", duration_ms=350, delay_ms=80)
    tl.add_effect(ui_card_3, "wipe", trigger="after_previous", direction="right", duration_ms=350)
    tl.add_effect(ui_card_1, "scale", trigger="with_previous", scale_factor=(1.06, 1.06), duration_ms=250, auto_reverse=True)
    tl.add_effect(footer, "fly_in", trigger="after_previous", direction="up", duration_ms=300)
    tl.add_effect(signal_orb, "fade_out", trigger="with_previous", duration_ms=250)

    tl.compile_to_openxml(slide)
    deck.save(output_path)
    return os.path.abspath(output_path)

