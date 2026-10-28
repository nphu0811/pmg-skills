"""
pmg.validator
=============
Deep OpenXML static validator and visibility-aware ExpectedTimeline extractor for PMG (v1.5.1).
Verifies:
- Valid ZIP / OpenXML package structure
- Unique shape identifiers per slide
- Valid transition XML and duration attributes
- Valid intra-slide <p:timing> animation trees, spTgt references, motion paths,
  <p:set> visibility transitions, grpId / <p:bldLst> synchronization, mainSeq click-group
  hierarchy, and conflicting duplicate entrance detection (v1.5.1 Upgrade D)
- Continuous '!!' Morph anchors across active sequences
- Valid media relationships (ppt/media/*)
- Native editable PowerPoint objects check (blocks flattened screenshot shortcuts)
- Object geometry, semantic roles, visibility windows, and ExpectedTimeline generation
"""

import os
import re
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from urllib.parse import unquote, urlsplit
from typing import Dict, List, Any, Tuple, Optional

from pptx import Presentation

from pmg.version import VERSION
from pmg.models import (
    ValidationStatus,
    ObjectMotionSpec,
    SlideTimelineSpec,
    ExpectedTimeline,
    ReviewDefect,
    AuditReport,
)
from pmg.animations import parse_slide_timing_openxml

EMU_PER_INCH = 914400.0

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
}


def inspect_package_integrity(package: zipfile.ZipFile) -> List[ReviewDefect]:
    """Read all OPC declarations, XML and relationships without writing a part."""
    defects: List[ReviewDefect] = []

    def defect(part, message, hint="repair_pptx_openxml"):
        defects.append(ReviewDefect("openxml_structure", "critical", 0, part, message, hint))

    names = set(package.namelist())
    for part, count in Counter(package.namelist()).items():
        if count > 1:
            defect(part, f"Duplicate package part '{part}' ({count} copies).")
    for part in ("[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"):
        if part not in names:
            defect(part, f"Missing required package part '{part}'.")
    xml_roots = {}
    for part in sorted(names):
        if part.endswith((".xml", ".rels")):
            try:
                xml_roots[part] = ET.fromstring(package.read(part))
            except ET.ParseError as exc:
                defect(part, f"Invalid XML in '{part}': {exc}")
    ct = xml_roots.get("[Content_Types].xml")
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    if ct is not None:
        if ct.tag != f"{{{ct_ns}}}Types":
            defect("[Content_Types].xml", "Invalid Content Types root element or namespace.")
        defaults, overrides = {}, {}
        for node in ct:
            if node.tag == f"{{{ct_ns}}}Default":
                key = node.get("Extension", "").lower()
                target = defaults
            elif node.tag == f"{{{ct_ns}}}Override":
                key = unquote(node.get("PartName", "")).lstrip("/")
                target = overrides
                if key not in names:
                    defect(key, f"ContentType override references missing part '{key}'.")
            else:
                defect("[Content_Types].xml", f"Unknown ContentType declaration '{node.tag}'.")
                continue
            mime = node.get("ContentType", "")
            if key in target or not key or "/" not in mime:
                defect(key, f"Duplicate or invalid ContentType declaration for '{key}'.")
            target[key] = mime
        for part in sorted(names):
            if part.endswith("/") or part == "[Content_Types].xml":
                continue
            extension = part.rsplit(".", 1)[-1].lower() if "." in part else ""
            mime = overrides.get(part, defaults.get(extension))
            if not mime:
                defect(part, f"Package part '{part}' has no registered ContentType in [Content_Types].xml.",
                       "declare_part_content_type")
            elif part == "ppt/pmg_execution_manifest.json" and mime != "application/json":
                defect(part, "Execution manifest must have ContentType application/json.", "declare_part_content_type")

    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    relationships_by_owner = {}
    for rel_path, root in xml_roots.items():
        if not rel_path.endswith(".rels"):
            continue
        if root.tag != f"{{{rel_ns}}}Relationships":
            defect(rel_path, f"Invalid Relationships namespace/root in '{rel_path}'.")
            continue
        if rel_path == "_rels/.rels":
            owner = ""
        elif "/_rels/" in rel_path:
            directory, filename = rel_path.rsplit("/_rels/", 1)
            owner = posixpath.join(directory, filename[:-5])
            if owner not in names:
                defect(rel_path, f"Relationships owner '{owner}' is missing.")
        else:
            defect(rel_path, f"Invalid relationships part location '{rel_path}'.")
            continue
        ids = set()
        for rel in root:
            rel_id, target = rel.get("Id", ""), rel.get("Target", "")
            if not rel_id or rel_id in ids or not rel.get("Type") or not target:
                defect(rel_path, f"Invalid or duplicate relationship '{rel_id}' in '{rel_path}'.")
            ids.add(rel_id)
            if rel.get("TargetMode") == "External":
                continue
            parsed = urlsplit(target)
            path = unquote(parsed.path)
            resolved = posixpath.normpath(path.lstrip("/") if path.startswith("/") else
                                         posixpath.join(posixpath.dirname(owner), path))
            if parsed.scheme or resolved.startswith("../") or resolved not in names:
                defect(rel_path, f"Relationship '{rel_id}' references missing/invalid package part '{resolved}'.")
        relationships_by_owner[owner] = ids
    for owner, owner_root in xml_roots.items():
        ids = relationships_by_owner.get(owner, set())
        for element in owner_root.iter():
            for attribute, reference in element.attrib.items():
                if attribute.startswith(f"{{{NS['r']}}}") and reference not in ids:
                    defect(owner, f"Object references undefined relationship '{reference}' in '{owner}'.")
    return defects


def inspect_powerpoint_compatibility(audit: Dict[str, Any], clean_open: Dict[str, Any]) -> List[ReviewDefect]:
    """Compare read-only COM observations with native XML effects."""
    if not clean_open.get("is_clean") or clean_open.get("skipped"):
        return [ReviewDefect("powerpoint_compatibility", "critical", 0, None,
                             clean_open.get("error", "Clean PowerPoint opening was not verified."))]
    defects = []
    for slide in audit.get("slides", []):
        number = slide["slide_number"]
        actual = clean_open.get("animation_counts", {}).get(str(number))
        if actual != slide["native_animation_count"]:
            defects.append(ReviewDefect("powerpoint_compatibility", "critical", number, None,
                                        f"PowerPoint recognized {actual} effects; XML declares {slide['native_animation_count']}.",
                                        "rebuild_timing_tree"))
    return defects


def _classify_semantic_role(
    name: str,
    obj_type: str,
    x: float,
    y: float,
    w: float,
    h: float,
    has_fill: bool,
    text: str,
) -> Tuple[str, bool]:
    """
    Returns (semantic_role, is_required_content) for visibility-aware validation.
    Explicitly separates decorative outlines (e.g. Oval rings with noFill) from hero images.
    """
    nl = (name or "").lower()
    tl = (text or "").lower()

    # 1. Full-slide backgrounds or tint scrims
    if (w >= 11.5 and h >= 6.5) or any(k in nl for k in ("background", "backdrop", "tint", "sky_backdrop", "starfield")):
        return ("background", False)

    # 2. Decorative outlines / rings (noFill shapes or Oval rings around hero objects)
    if not has_fill and obj_type == "shape" and not text:
        return ("decorative_outline", False)
    if any(k in nl for k in ("outline", "ring", "halo_ring", "decorative", "orbit_ring")) or (
        nl.startswith("oval") and not text and w >= 2.5 and h >= 2.5
    ):
        return ("decorative_outline", False)

    # 3. Titles & Subtitles (evaluated before hero keywords so '!!HeroTitle' / '!!HeroSubtitle' classify as title/subtitle)
    if obj_type != "picture":
        if "subtitle" in nl or ("sub" in nl and text):
            return ("subtitle", True)
        if "title" in nl or "herotitle" in nl or (text and y < 2.3 and w >= 3.2 and h >= 0.45 and len(text) >= 6):
            return ("title", True)

    # 4. Hero images & primary visual subjects
    hero_keywords = (
        "earth", "globe", "hero", "planet", "mars", "galaxy", "milky",
        "astronaut", "spacecraft", "station", "iss", "black_hole", "blackhole",
        "satellite", "solar_system", "nebula", "telescope",
    )
    if obj_type == "picture" and w >= 1.2 and h >= 1.2:
        return ("hero_image", True)
    if any(k in nl for k in hero_keywords) and w >= 1.0 and h >= 1.0:
        return ("hero_image", True)

    # 5. Diagrams & Data visualizations
    if obj_type == "diagram" or any(k in nl for k in ("diagram", "chart", "infographic", "bar_", "metric_bar")):
        return ("diagram", True)

    # 6. Information / UI Cards
    if any(k in nl for k in ("card", "widget", "panel", "module")) and w >= 1.4 and h >= 0.7:
        return ("card", True)
    if text and w >= 1.8 and h >= 0.55 and y >= 2.0:
        return ("card", True)

    return ("content", False)


def _parse_shape_node(elem: ET.Element, obj_type: str) -> Optional[Dict[str, Any]]:
    """Extracts id, name, position, size, rotation, alpha, fill state, text, and semantic role."""
    cNvPr = elem.find(".//p:cNvPr", NS)
    if cNvPr is None:
        return None
    shp_id = cNvPr.attrib.get("id", "0")
    name = cNvPr.attrib.get("name", "")

    xfrm = elem.find(".//a:xfrm", NS)
    x_in, y_in, w_in, h_in, rot_deg = 0.0, 0.0, 0.0, 0.0, 0.0
    if xfrm is not None:
        rot_str = xfrm.attrib.get("rot", "0")
        try:
            rot_deg = int(rot_str) / 60000.0
        except ValueError:
            rot_deg = 0.0
        off = xfrm.find("a:off", NS)
        ext = xfrm.find("a:ext", NS)
        if off is not None:
            x_in = int(off.attrib.get("x", "0")) / EMU_PER_INCH
            y_in = int(off.attrib.get("y", "0")) / EMU_PER_INCH
        if ext is not None:
            w_in = int(ext.attrib.get("cx", "0")) / EMU_PER_INCH
            h_in = int(ext.attrib.get("cy", "0")) / EMU_PER_INCH

    no_fill = elem.find("./p:spPr/a:noFill", NS) is not None
    has_fill = (obj_type == "picture") or (not no_fill)

    alpha_elem = elem.find(".//a:solidFill//a:alpha", NS)
    alpha_val = 1.0
    if alpha_elem is not None:
        try:
            alpha_val = int(alpha_elem.attrib.get("val", "100000")) / 100000.0
        except ValueError:
            alpha_val = 1.0

    texts = [t.text for t in elem.findall(".//a:t", NS) if t.text]
    full_text = " ".join(texts).strip()

    role, is_req = _classify_semantic_role(name, obj_type, x_in, y_in, w_in, h_in, has_fill, full_text)

    return {
        "id": shp_id,
        "name": name,
        "type": obj_type,
        "x": round(x_in, 4),
        "y": round(y_in, 4),
        "w": round(w_in, 4),
        "h": round(h_in, 4),
        "rot_deg": round(rot_deg, 2),
        "alpha": round(alpha_val, 3),
        "has_fill": has_fill,
        "text": full_text,
        "semantic_role": role,
        "is_required_content": is_req,
    }


def extract_slide_objects(xml_bytes: bytes) -> List[Dict[str, Any]]:
    """Parses all top-level and grouped shapes from slide XML."""
    root = ET.fromstring(xml_bytes)
    spTree = root.find(".//p:cSld/p:spTree", NS)
    if spTree is None:
        return []

    objects: List[Dict[str, Any]] = []
    tag_map = {
        f"{{{NS['p']}}}sp": "shape",
        f"{{{NS['p']}}}pic": "picture",
        f"{{{NS['p']}}}grpSp": "group",
        f"{{{NS['p']}}}cxnSp": "connector",
        f"{{{NS['p']}}}graphicFrame": "diagram",
    }

    def visit(parent, parent_id=None):
        for z_order, child in enumerate(parent):
            obj_type = tag_map.get(child.tag)
            if not obj_type:
                continue
            parsed = _parse_shape_node(child, obj_type)
            if parsed:
                parsed.update(parent_id=parent_id, z_order=z_order)
                if obj_type == "shape" and parsed["text"]:
                    parsed["type"] = "textbox" if "TextBox" in parsed["name"] else "shape"
                if obj_type == "group":
                    parsed["text"] = ""  # Children are inspected separately.
                objects.append(parsed)
                if obj_type == "group":
                    visit(child, parsed["id"])

    visit(spTree)
    return objects


def parse_slide_transition(xml_text: str, slide_idx: int) -> Tuple[str, int, bool]:
    """Returns (transition_type, duration_ms, has_explicit_transition)."""
    trans_match = re.search(r"<p:transition([^>]*)>(.*?)</p:transition>", xml_text, re.DOTALL)
    if not trans_match:
        return ("none" if slide_idx == 1 else "missing", 0, False)

    t_attrs = trans_match.group(1)
    t_inner = trans_match.group(2)
    dur_m = re.search(r'dur="(\d+)"', t_attrs)
    dur_ms = int(dur_m.group(1)) if dur_m else 1600

    if "morph" in t_inner:
        return ("morph", dur_ms, True)
    elif "<p:circle" in t_inner:
        return ("circle", dur_ms, True)
    elif "<p:split" in t_inner:
        return ("split", dur_ms, True)
    elif "<p:fade" in t_inner:
        return ("fade", dur_ms, True)
    elif "<p:push" in t_inner:
        return ("push", dur_ms, True)
    elif "<p:wipe" in t_inner:
        return ("wipe", dur_ms, True)
    else:
        return ("native", dur_ms, True)


def extract_expected_timeline(pptx_path: str, hold_duration_ms: int = 2000) -> ExpectedTimeline:
    """
    Constructs a deterministic, visibility-aware ExpectedTimeline from a PPTX file, capturing:
    - Inter-slide Morph object movements (initial/final pos, size, rotation, opacity)
    - Intra-slide native animations (in <p:timing>)
    - Expected initial visibility, appearance time, disappearance time, and completion state
    - Absolute, slide-relative, and click-triggered playback script
    """
    deck_title = os.path.splitext(os.path.basename(pptx_path))[0]
    slides_spec: List[SlideTimelineSpec] = []
    playback_script: List[Dict[str, Any]] = []

    prs = Presentation(pptx_path)

    with zipfile.ZipFile(pptx_path, "r") as z:
        slide_files = [n for n in z.namelist() if re.match(r"^ppt/slides/slide\d+\.xml$", n)]
        slide_files.sort(key=lambda x: int(re.search(r"\d+", x).group()))

        prev_objects_by_name: Dict[str, Dict[str, Any]] = {}
        current_time_ms = 0

        playback_script.append({
            "event": "start_slideshow",
            "slide_number": 1,
            "time_ms": 0,
        })

        for idx, s_file in enumerate(slide_files, 1):
            xml_bytes = z.read(s_file)
            xml_text = xml_bytes.decode("utf-8", errors="ignore")
            objects = extract_slide_objects(xml_bytes)
            trans_type, trans_dur, _ = parse_slide_transition(xml_text, idx)
            if idx == 1:
                trans_dur = 0

            slide_start_ms = current_time_ms
            motion_specs: List[ObjectMotionSpec] = []
            seq_idx = 0

            curr_objects_by_name = {o["name"]: o for o in objects if o["name"]}
            static_spec_by_name: Dict[str, ObjectMotionSpec] = {}

            # 1. Track Morph / Cross-slide object movements and baseline static specs
            for obj in objects:
                name = obj["name"] or f"shape_{obj['id']}"
                role = obj.get("semantic_role", "content")
                is_req = obj.get("is_required_content", False)

                if idx > 1 and name.startswith("!!") and name in prev_objects_by_name:
                    prev_obj = prev_objects_by_name[name]
                    seq_idx += 1
                    spec = ObjectMotionSpec(
                        object_id=name,
                        slide_number=idx,
                        object_type=obj["type"],
                        initial_position=(prev_obj["x"], prev_obj["y"]),
                        final_position=(obj["x"], obj["y"]),
                        initial_size=(prev_obj["w"], prev_obj["h"]),
                        final_size=(obj["w"], obj["h"]),
                        rotation_params={
                            "start_deg": prev_obj["rot_deg"],
                            "end_deg": obj["rot_deg"],
                        },
                        opacity_params={
                            "start_alpha": prev_obj["alpha"],
                            "end_alpha": obj["alpha"],
                        },
                        motion_path=f"morph_linear:({prev_obj['x']},{prev_obj['y']})->({obj['x']},{obj['y']})",
                        start_time_ms=0,
                        duration_ms=trans_dur if trans_dur > 0 else 1600,
                        delay_ms=0,
                        trigger_type="slide_transition",
                        animation_sequence=seq_idx,
                        easing="smooth_morph",
                        time_reference="slide_relative",
                        absolute_start_ms=slide_start_ms,
                        absolute_end_ms=slide_start_ms + (trans_dur if trans_dur > 0 else 1600),
                        semantic_role=role,
                        is_required_content=is_req,
                        expected_initial_visibility=True,
                        expected_appearance_time_ms=0,
                        expected_disappearance_time_ms=None,
                        expected_completion_state="visible",
                    )
                    motion_specs.append(spec)
                    static_spec_by_name[name] = spec
                else:
                    spec = ObjectMotionSpec(
                        object_id=name,
                        slide_number=idx,
                        object_type=obj["type"],
                        initial_position=(obj["x"], obj["y"]),
                        final_position=(obj["x"], obj["y"]),
                        initial_size=(obj["w"], obj["h"]),
                        final_size=(obj["w"], obj["h"]),
                        rotation_params={"start_deg": obj["rot_deg"], "end_deg": obj["rot_deg"]},
                        opacity_params={"start_alpha": obj["alpha"], "end_alpha": obj["alpha"]},
                        motion_path=None,
                        start_time_ms=0,
                        duration_ms=0,
                        delay_ms=0,
                        trigger_type="static",
                        animation_sequence=0,
                        easing="none",
                        time_reference="slide_relative",
                        absolute_start_ms=slide_start_ms,
                        absolute_end_ms=slide_start_ms + trans_dur + hold_duration_ms,
                        semantic_role=role,
                        is_required_content=is_req,
                        expected_initial_visibility=True,
                        expected_appearance_time_ms=0,
                        expected_disappearance_time_ms=None,
                        expected_completion_state="visible",
                    )
                    motion_specs.append(spec)
                    static_spec_by_name[name] = spec

            # 2. Parse intra-slide <p:timing> animations using pmg.animations
            clicks: List[Dict[str, Any]] = []
            max_anim_end_rel = trans_dur
            if idx <= len(prs.slides):
                intra_specs, intra_clicks = parse_slide_timing_openxml(
                    prs.slides[idx - 1],
                    slide_number=idx,
                    slide_abs_start_ms=slide_start_ms,
                    transition_offset_ms=trans_dur,
                )
                for ispec in intra_specs:
                    seq_idx += 1
                    ispec.animation_sequence = seq_idx
                    base_obj = static_spec_by_name.get(ispec.object_id)
                    if base_obj is not None:
                        ispec.semantic_role = base_obj.semantic_role
                        ispec.is_required_content = base_obj.is_required_content
                        ispec.object_type = base_obj.object_type

                    is_entr = ispec.easing.startswith("entr_") or ispec.easing in (
                        "appear", "fade", "fly_in", "wipe", "zoom"
                    )
                    is_exit = ispec.easing.startswith("exit_") or ispec.easing in (
                        "disappear", "fade_out", "fly_out", "wipe_out"
                    )
                    if is_entr:
                        ispec.expected_initial_visibility = False
                        ispec.expected_appearance_time_ms = ispec.start_time_ms + ispec.duration_ms
                        ispec.expected_completion_state = "visible"
                        if base_obj is not None:
                            base_obj.expected_initial_visibility = False
                            base_obj.expected_appearance_time_ms = ispec.expected_appearance_time_ms
                            base_obj.expected_completion_state = "visible"
                    elif is_exit:
                        ispec.expected_disappearance_time_ms = ispec.start_time_ms + ispec.duration_ms
                        ispec.expected_completion_state = "hidden"
                        if base_obj is not None:
                            base_obj.expected_disappearance_time_ms = ispec.expected_disappearance_time_ms
                            base_obj.expected_completion_state = "hidden"

                    motion_specs.append(ispec)
                    rel_end = ispec.absolute_end_ms - slide_start_ms
                    if rel_end > max_anim_end_rel:
                        max_anim_end_rel = rel_end
                clicks.extend(intra_clicks)

            slide_end_ms = slide_start_ms + max_anim_end_rel + hold_duration_ms

            if idx > 1:
                playback_script.append({
                    "event": "advance_slide",
                    "from_slide": idx - 1,
                    "to_slide": idx,
                    "time_ms": slide_start_ms,
                    "transition_type": trans_type,
                    "transition_duration_ms": trans_dur,
                })

            for cl in clicks:
                playback_script.append({
                    "event": "click_trigger",
                    "slide_number": idx,
                    "sequence_index": cl["sequence_index"],
                    "target_object": cl["target_object"],
                    "time_ms": cl["absolute_ms"],
                })

            playback_script.append({
                "event": "hold_slide",
                "slide_number": idx,
                "start_ms": slide_start_ms + max_anim_end_rel,
                "end_ms": slide_end_ms,
                "hold_duration_ms": hold_duration_ms,
            })

            slides_spec.append(
                SlideTimelineSpec(
                    slide_number=idx,
                    transition_type=trans_type,
                    transition_duration_ms=trans_dur,
                    hold_duration_ms=hold_duration_ms,
                    absolute_start_ms=slide_start_ms,
                    absolute_end_ms=slide_end_ms,
                    objects=motion_specs,
                    clicks=clicks,
                )
            )

            prev_objects_by_name = curr_objects_by_name
            current_time_ms = slide_end_ms

        playback_script.append({
            "event": "end_slideshow",
            "time_ms": current_time_ms,
        })

    return ExpectedTimeline(
        deck_title=deck_title,
        engine_version=VERSION,
        total_slides=len(slides_spec),
        total_duration_ms=current_time_ms,
        slides=slides_spec,
        playback_script=playback_script,
    )


def verify_clean_powerpoint_open(pptx_path: str, timeout_sec: float = 30.0) -> Dict[str, Any]:
    """Backward compatible facade for the shared read-only worker probe."""
    from pmg.worker.powerpoint_probe import verify_clean_powerpoint_open as probe
    return probe(pptx_path, timeout_sec=timeout_sec)


class StaticValidator:
    """
    Performs static OpenXML, geometry, media, transition, intra-slide animation timing,
    disappearing-object root-cause detection, and anti-shortcut validation.
    """

    def validate(self, pptx_path: str, verify_com_open: bool = False) -> Dict[str, Any]:
        defects: List[ReviewDefect] = []
        slide_summaries: List[Dict[str, Any]] = []
        traveling_anchors: Dict[str, List[int]] = {}
        discontinuous_anchors: Dict[str, List[int]] = {}
        missing_transitions: List[int] = []
        total_native_animations = 0

        if not os.path.exists(pptx_path):
            defects.append(
                ReviewDefect(
                    category="openxml_structure",
                    severity="critical",
                    slide_number=0,
                    object_id=None,
                    description=f"Presentation file does not exist: {pptx_path}",
                )
            )
            return AuditReport(defects, {"total_slides": 0}).to_dict()

        try:
            with zipfile.ZipFile(pptx_path, "r") as z:
                namelist = set(z.namelist())
                defects.extend(inspect_package_integrity(z))

                if "ppt/presentation.xml" not in namelist:
                    defects.append(
                        ReviewDefect(
                            category="openxml_structure",
                            severity="critical",
                            slide_number=0,
                            object_id=None,
                            description="Missing ppt/presentation.xml in package.",
                        )
                    )
                    return AuditReport(defects, {"total_slides": 0}).to_dict()

                pres_xml = z.read("ppt/presentation.xml").decode("utf-8", errors="ignore")
                sldSz = re.search(r'<p:sldSz[^>]*cx="(\d+)"[^>]*cy="(\d+)"', pres_xml)
                if sldSz:
                    cx, cy = int(sldSz.group(1)), int(sldSz.group(2))
                    ratio = cx / cy if cy > 0 else 0
                    if abs(ratio - (16.0 / 9.0)) > 0.05:
                        defects.append(
                            ReviewDefect(
                                category="visual_quality",
                                severity="major",
                                slide_number=0,
                                object_id=None,
                                description=f"Canvas aspect ratio is not 16:9 (cx={cx}, cy={cy}).",
                            )
                        )

                slide_files = [n for n in namelist if re.match(r"^ppt/slides/slide\d+\.xml$", n)]
                slide_files.sort(key=lambda x: int(re.search(r"\d+", x).group()))
                total_slides = len(slide_files)

                if total_slides == 0:
                    defects.append(
                        ReviewDefect(
                            category="openxml_structure",
                            severity="critical",
                            slide_number=0,
                            object_id=None,
                            description="Presentation contains 0 slides.",
                        )
                    )

                for idx, s_file in enumerate(slide_files, 1):
                    xml_bytes = z.read(s_file)
                    xml_text = xml_bytes.decode("utf-8", errors="ignore")
                    objects = extract_slide_objects(xml_bytes)

                    seen_ids = set()
                    obj_name_by_id: Dict[str, str] = {}
                    for o in objects:
                        obj_name_by_id[o["id"]] = o["name"] or f"spid_{o['id']}"
                        if o["id"] in seen_ids:
                            defects.append(
                                ReviewDefect(
                                    category="openxml_structure",
                                    severity="major",
                                    slide_number=idx,
                                    object_id=o["name"],
                                    description=f"Duplicate shape ID {o['id']} on slide {idx}.",
                                )
                            )
                        seen_ids.add(o["id"])

                    # 2. Check Native Editable Objects (Anti-Shortcut Rule)
                    editable_objs = [o for o in objects if o["type"] in ("shape", "textbox", "group", "diagram", "connector")]
                    pictures = [o for o in objects if o["type"] == "picture"]
                    if len(editable_objs) == 0 and len(pictures) >= 1:
                        defects.append(
                            ReviewDefect(
                                category="openxml_structure",
                                severity="critical",
                                slide_number=idx,
                                object_id=pictures[0]["name"],
                                description=f"Slide {idx} contains only image(s) and zero editable shapes/text (flattened slide shortcut is forbidden).",
                            )
                        )
                    elif len(objects) < 2:
                        defects.append(
                            ReviewDefect(
                                category="visual_quality",
                                severity="major",
                                slide_number=idx,
                                object_id=None,
                                description=f"Slide {idx} has fewer than 2 independent objects ({len(objects)} found).",
                            )
                        )

                    # 3. Check Transition
                    trans_type, trans_dur, has_trans = parse_slide_transition(xml_text, idx)
                    if idx > 1 and not has_trans:
                        missing_transitions.append(idx)
                        defects.append(
                            ReviewDefect(
                                category="animation_behavior",
                                severity="major",
                                slide_number=idx,
                                object_id=None,
                                description=f"Slide {idx} is missing a native PowerPoint transition.",
                                remediation_hint="inject_transition",
                            )
                        )
                    elif idx > 1 and (trans_dur < 300 or trans_dur > 5000):
                        defects.append(
                            ReviewDefect(
                                category="temporal_quality",
                                severity="minor",
                                slide_number=idx,
                                object_id=None,
                                description=f"Slide {idx} transition duration ({trans_dur}ms) is outside recommended bounds (300ms-5000ms).",
                                remediation_hint="normalize_duration",
                            )
                        )

                    # 4. Validate Intra-Slide <p:timing> Native Animations & Disappearing Object Root Causes
                    root = ET.fromstring(xml_bytes)
                    timing = root.find("p:timing", NS)
                    slide_anim_count = 0
                    if timing is not None:
                        main_seq = timing.find('.//p:cTn[@nodeType="mainSeq"]', NS)
                        if main_seq is None:
                            defects.append(
                                ReviewDefect(
                                    category="openxml_structure",
                                    severity="major",
                                    slide_number=idx,
                                    object_id=None,
                                    description=f"Slide {idx} has <p:timing> without a valid mainSeq node.",
                                    remediation_hint="rebuild_timing_tree",
                                )
                            )
                        else:
                            # Root Cause Check C: Multiple top-level <p:par> under mainSeq with withEffect
                            # or missing <p:tn> in onBegin condition
                            ms_child_lst = main_seq.find("p:childTnLst", NS)
                            if ms_child_lst is not None:
                                top_pars = ms_child_lst.findall("p:par", NS)
                                if top_pars:
                                    first_ctn = top_pars[0].find("p:cTn", NS)
                                    if first_ctn is not None:
                                        on_begin = first_ctn.find('.//p:stCondLst/p:cond[@evt="onBegin"]', NS)
                                        if on_begin is not None and on_begin.find("p:tn", NS) is None:
                                            defects.append(
                                                ReviewDefect(
                                                    category="object_visibility",
                                                    severity="critical",
                                                    slide_number=idx,
                                                    object_id=None,
                                                    description=(
                                                        f"Slide {idx} mainSeq click-group <p:cond evt='onBegin'> is missing "
                                                        f"<p:tn val='2'/> target reference, preventing reliable auto-start in Slide Show."
                                                    ),
                                                    remediation_hint="repair_disappearing_object_timing",
                                                )
                                            )
                                    if len(top_pars) > 1:
                                        for sec_par in top_pars[1:]:
                                            sec_eff = sec_par.find('.//p:cTn[@presetClass]', NS)
                                            if sec_eff is not None and sec_eff.attrib.get("nodeType") == "withEffect":
                                                sp_t = sec_eff.find(".//p:spTgt", NS)
                                                t_id = sp_t.attrib.get("spid", "") if sp_t is not None else ""
                                                t_name = obj_name_by_id.get(t_id, f"spid_{t_id}")
                                                defects.append(
                                                    ReviewDefect(
                                                        category="object_visibility",
                                                        severity="critical",
                                                        slide_number=idx,
                                                        object_id=t_name,
                                                        description=(
                                                            f"Slide {idx} places 'withEffect' animation for '{t_name}' in a separate "
                                                            f"top-level <p:seq> child <p:par>, causing '{t_name}' to remain hidden in Slide Show."
                                                        ),
                                                        remediation_hint="repair_disappearing_object_timing",
                                                    )
                                                )

                        # Gather <p:bldLst> entries
                        bld_grps_by_spid: Dict[str, List[str]] = {}
                        for bld_p in timing.findall(".//p:bldLst/p:bldP", NS):
                            b_spid = bld_p.attrib.get("spid", "")
                            b_grp = bld_p.attrib.get("grpId", "")
                            bld_grps_by_spid.setdefault(b_spid, []).append(b_grp)

                        active_visible_state: Dict[str, str] = {}  # spid -> last presetClass ("entr" or "exit")
                        defined_grp_ids = set()

                        for ctn in timing.findall(".//p:cTn", NS):
                            p_cls = ctn.attrib.get("presetClass")
                            if p_cls in ("entr", "emph", "path", "exit"):
                                slide_anim_count += 1
                                grp_id_attr = ctn.attrib.get("grpId")
                                if grp_id_attr is not None:
                                    defined_grp_ids.add(grp_id_attr)

                                sp_tgts = ctn.findall(".//p:spTgt", NS)
                                if not sp_tgts:
                                    defects.append(
                                        ReviewDefect(
                                            category="animation_behavior",
                                            severity="major",
                                            slide_number=idx,
                                            object_id=None,
                                            description=f"Animation effect on slide {idx} is missing a <p:spTgt> target.",
                                            remediation_hint="repair_animation_target",
                                        )
                                    )
                                else:
                                    tgt_spid = sp_tgts[0].attrib.get("spid", "")
                                    tgt_name = obj_name_by_id.get(tgt_spid, f"spid_{tgt_spid}")
                                    for st in sp_tgts:
                                        s_id = st.attrib.get("spid", "")
                                        if s_id not in seen_ids:
                                            defects.append(
                                                ReviewDefect(
                                                    category="animation_behavior",
                                                    severity="major",
                                                    slide_number=idx,
                                                    object_id=f"spid_{s_id}",
                                                    description=f"Animation effect on slide {idx} targets nonexistent shape spid={s_id}.",
                                                    remediation_hint="repair_animation_target",
                                                )
                                            )

                                    # Root Cause Check A: Entrance animation missing <p:set> style.visibility="visible"
                                    if p_cls == "entr":
                                        has_vis_set = False
                                        for p_set in ctn.findall(".//p:set", NS):
                                            attr_el = p_set.find(".//p:attrName", NS)
                                            val_el = p_set.find(".//p:to/p:strVal", NS)
                                            if (
                                                attr_el is not None
                                                and (attr_el.text or "").strip() == "style.visibility"
                                                and val_el is not None
                                                and val_el.attrib.get("val") == "visible"
                                            ):
                                                has_vis_set = True
                                                break
                                        if not has_vis_set:
                                            defects.append(
                                                ReviewDefect(
                                                    category="object_visibility",
                                                    severity="critical",
                                                    slide_number=idx,
                                                    object_id=tgt_name,
                                                    description=(
                                                        f"Entrance animation on '{tgt_name}' (slide {idx}) is missing "
                                                        f"<p:set> style.visibility='visible', causing '{tgt_name}' to disappear during Slide Show."
                                                    ),
                                                    remediation_hint="repair_disappearing_object_timing",
                                                )
                                            )

                                        # Root Cause Check B: Missing grpId when spid is in <p:bldLst>
                                        if tgt_spid in bld_grps_by_spid and grp_id_attr is None:
                                            defects.append(
                                                ReviewDefect(
                                                    category="object_visibility",
                                                    severity="critical",
                                                    slide_number=idx,
                                                    object_id=tgt_name,
                                                    description=(
                                                        f"Entrance animation on '{tgt_name}' (slide {idx}) omits grpId while "
                                                        f"<p:bldLst> binds spid={tgt_spid}, leaving '{tgt_name}' hidden in Slide Show."
                                                    ),
                                                    remediation_hint="repair_disappearing_object_timing",
                                                )
                                            )

                                        # Root Cause Check D: Duplicate conflicting entrance animations on same spid
                                        if active_visible_state.get(tgt_spid) == "entr":
                                            defects.append(
                                                ReviewDefect(
                                                    category="object_visibility",
                                                    severity="major",
                                                    slide_number=idx,
                                                    object_id=tgt_name,
                                                    description=(
                                                        f"Duplicate conflicting entrance animation on '{tgt_name}' (slide {idx}) "
                                                        f"without an intervening exit effect."
                                                    ),
                                                    remediation_hint="repair_disappearing_object_timing",
                                                )
                                            )
                                        active_visible_state[tgt_spid] = "entr"
                                    elif p_cls == "exit":
                                        active_visible_state[tgt_spid] = "exit"

                                for am in ctn.findall(".//p:animMotion", NS):
                                    path_attr = am.attrib.get("path", "").strip()
                                    if not path_attr or not path_attr.upper().startswith("M"):
                                        defects.append(
                                            ReviewDefect(
                                                category="animation_behavior",
                                                severity="major",
                                                slide_number=idx,
                                                object_id=None,
                                                description=f"Invalid motion path syntax '{path_attr}' on slide {idx}.",
                                                remediation_hint="repair_motion_path",
                                            )
                                        )

                    total_native_animations += slide_anim_count

                    # 5. Track !! Morph anchors
                    for o in objects:
                        if o["name"].startswith("!!"):
                            traveling_anchors.setdefault(o["name"], []).append(idx)

                    # 6. Check Media Relationships
                    rels_path = f"ppt/slides/_rels/slide{idx}.xml.rels"
                    if rels_path in namelist:
                        rels_xml = z.read(rels_path).decode("utf-8", errors="ignore")
                        targets = re.findall(r'Target="\.\./media/([^"]+)"', rels_xml)
                        for tgt in targets:
                            media_file = f"ppt/media/{tgt}"
                            if media_file not in namelist:
                                defects.append(
                                    ReviewDefect(
                                        category="openxml_structure",
                                        severity="critical",
                                        slide_number=idx,
                                        object_id=tgt,
                                        description=f"Slide {idx} references missing media file {media_file}.",
                                    )
                                )

                    word_count = sum(len(o["text"].split()) for o in objects if o["text"])

                    slide_summaries.append({
                        "slide_number": idx,
                        "object_count": len(objects),
                        "editable_object_count": len(editable_objs),
                        "picture_count": len(pictures),
                        "required_content_objects": [o["name"] for o in objects if o.get("is_required_content")],
                        "decorative_outlines": [o["name"] for o in objects if o.get("semantic_role") == "decorative_outline"],
                        "transition_type": trans_type,
                        "transition_duration_ms": trans_dur,
                        "native_animation_count": slide_anim_count,
                        "word_count": word_count,
                        "morph_anchors": [o["name"] for o in objects if o["name"].startswith("!!")],
                    })

                for tag, s_list in traveling_anchors.items():
                    s_list.sort()
                    if len(s_list) > 1 and ("Traveling" in tag or "prog_" in tag):
                        gaps = []
                        for s_num in range(min(s_list), max(s_list) + 1):
                            if s_num not in s_list:
                                gaps.append(s_num)
                        if gaps:
                            discontinuous_anchors[tag] = gaps
                            defects.append(
                                ReviewDefect(
                                    category="animation_behavior",
                                    severity="major",
                                    slide_number=gaps[0],
                                    object_id=tag,
                                    description=f"Traveling entity '{tag}' has continuity gap on slide(s) {gaps} (present on {s_list}).",
                                    remediation_hint="bridge_anchor_gap",
                                )
                            )

        except Exception as e:
            defects.append(
                ReviewDefect(
                    category="openxml_structure",
                    severity="critical",
                    slide_number=0,
                    object_id=None,
                    description=f"Failed to parse PPTX package: {str(e)}",
                )
            )
            return AuditReport(defects, {"total_slides": 0}).to_dict()

        if verify_com_open and not any(d.severity == "critical" for d in defects):
            clean_res = verify_clean_powerpoint_open(pptx_path)
            if not clean_res.get("is_clean", False):
                defects.append(
                    ReviewDefect(
                        category="openxml_structure",
                        severity="critical",
                        slide_number=0,
                        object_id=None,
                        description=clean_res.get("error", "PowerPoint Repair or Open failure"),
                        remediation_hint="repair_pptx_openxml",
                    )
                )

        return AuditReport(defects=defects, details={
            "total_slides": len(slide_summaries),
            "total_native_animations": total_native_animations,
            "missing_transitions": missing_transitions,
            "traveling_anchors": traveling_anchors,
            "discontinuous_anchors": discontinuous_anchors,
            "slides": slide_summaries,
        }).to_dict()
