"""Acceptance, Regression, and Integration Tests for PMG Skills v1.8.0:
- Advanced Cinematic Compositing (Massive hierarchies, Bezier/orbital paths, multi-plane parallax)
- Expanded 44-Case Animation Benchmark Suite (Groups A through H)
- Mandatory Acceptance Decks:
  1. Agribank Motion Identity System (8 slides)
  2. Cinematic Shutter & Perspective Reveal (5 slides)
  3. Deep Space Exploration — Beyond the Horizon (12 slides, >100 objects)
- Microsoft PowerPoint Desktop COM Clean Open Verification (DisplayAlerts=2, OpenAndRepair=False)
"""
import os
import pytest
from pathlib import Path

from pmg.version import VERSION
from pmg.validator import StaticValidator, verify_clean_powerpoint_open
from pmg.compositing import (
    build_bezier_motion_path,
    build_orbital_motion_path,
    build_waypoint_motion_path,
    generate_massive_constellation,
    build_v1_8_agribank_identity_deck,
    build_v1_8_cinematic_shutter_reveal,
    build_v1_8_deep_space_exploration,
    PARALLAX_PLANES,
)
from benchmarks.cases import BENCHMARK_CASES, get_benchmark_case


def test_v1_8_version():
    assert VERSION in ("1.8.0", "2.0.1")


def test_motion_path_generators():
    # 1. Cubic Bezier path
    bez = build_bezier_motion_path((0.1, 0.2), (0.3, 0.4), (0.6, 0.8), (0.9, 0.95))
    assert bez.startswith("M 0.1000 0.2000 C")
    assert bez.endswith("E")

    # 2. Orbital path
    orb = build_orbital_motion_path(cx=0.5, cy=0.5, rx=0.2, ry=0.1)
    assert orb.startswith("M")
    assert "L" in orb
    assert orb.endswith("E")

    # 3. Waypoint path
    way = build_waypoint_motion_path([(0.0, 0.0), (0.2, -0.1), (0.4, 0.1), (0.6, 0.0)])
    assert way == "M 0.0000 0.0000 L 0.2000 -0.1000 L 0.4000 0.1000 L 0.6000 0.0000 E"


def test_massive_scene_hierarchy_generation():
    nodes_50 = generate_massive_constellation(num_nodes=50, seed=1)
    assert len(nodes_50) == 50

    nodes_100 = generate_massive_constellation(num_nodes=100, seed=2)
    assert len(nodes_100) == 100

    nodes_200 = generate_massive_constellation(num_nodes=200, seed=3)
    assert len(nodes_200) == 200

    for n in nodes_100:
        assert 0.5 <= n["x"] <= 12.8
        assert 1.5 <= n["y"] <= 6.8
        assert n["size"] > 0


def test_44_case_benchmark_suite_registration():
    assert len(BENCHMARK_CASES) == 44
    groups = {c.group for c in BENCHMARK_CASES}
    assert groups == {"Group A", "Group B", "Group C", "Group D", "Group E", "Group F", "Group G", "Group H"}

    # Test retrieving first, middle, boundary, and last cases
    c1 = get_benchmark_case(1)
    assert c1.slug == "case_01_fade_in"
    c24 = get_benchmark_case(24)
    assert c24.slug == "case_24_failure_detection_and_repair"
    c25 = get_benchmark_case(25)
    assert c25.group == "Group E"
    c44 = get_benchmark_case(44)
    assert c44.group == "Group H"

    with pytest.raises(KeyError):
        get_benchmark_case(45)


def test_benchmark_cases_group_e_f_g_h_builders(tmp_path):
    # Test sample builders from each new group
    out_dir = str(tmp_path)

    # Group E: Case 25
    c25 = get_benchmark_case(25)
    p25 = os.path.join(out_dir, "test_case_25.pptx")
    c25.builder_fn(p25, out_dir)
    assert os.path.exists(p25)
    val25 = StaticValidator().validate(p25)
    assert val25["is_valid"] is True

    # Group F: Case 31 (100 nodes)
    c31 = get_benchmark_case(31)
    p31 = os.path.join(out_dir, "test_case_31.pptx")
    c31.builder_fn(p31, out_dir)
    assert os.path.exists(p31)
    val31 = StaticValidator().validate(p31)
    assert val31["is_valid"] is True
    import pptx
    prs31 = pptx.Presentation(p31)
    assert len(prs31.slides[0].shapes) >= 100

    # Group G: Case 36 (10-strip shutter)
    c36 = get_benchmark_case(36)
    p36 = os.path.join(out_dir, "test_case_36.pptx")
    c36.builder_fn(p36, out_dir)
    assert os.path.exists(p36)
    val36 = StaticValidator().validate(p36)
    assert val36["is_valid"] is True

    # Group H: Case 44 (clean open com verification)
    c44 = get_benchmark_case(44)
    p44 = os.path.join(out_dir, "test_case_44.pptx")
    c44.builder_fn(p44, out_dir)
    assert os.path.exists(p44)
    clean44 = verify_clean_powerpoint_open(p44)
    assert clean44["is_clean"] is True


def test_acceptance_demo1_agribank_identity_deck(tmp_path):
    out_pptx = str(tmp_path / "agribank_v1_8_identity.pptx")
    deck_path = build_v1_8_agribank_identity_deck(out_pptx)
    assert os.path.exists(deck_path)

    val = StaticValidator().validate(deck_path)
    assert val["is_valid"] is True, f"Defects: {val.get('defects')}"
    assert len(val["defects"]) == 0
    assert val["total_slides"] == 8

    # Clean COM open check
    clean = verify_clean_powerpoint_open(deck_path)
    assert clean["is_clean"] is True, f"COM Clean open error: {clean.get('error')}"


def test_acceptance_demo2_shutter_perspective_deck(tmp_path):
    out_pptx = str(tmp_path / "shutter_perspective_v1_8.pptx")
    deck_path = build_v1_8_cinematic_shutter_reveal(out_pptx)
    assert os.path.exists(deck_path)

    val = StaticValidator().validate(deck_path)
    assert val["is_valid"] is True, f"Defects: {val.get('defects')}"
    assert len(val["defects"]) == 0
    assert val["total_slides"] == 5

    clean = verify_clean_powerpoint_open(deck_path)
    assert clean["is_clean"] is True, f"COM Clean open error: {clean.get('error')}"


def test_acceptance_demo3_deep_space_exploration_deck(tmp_path):
    out_pptx = str(tmp_path / "deep_space_v1_8_exploration.pptx")
    deck_path = build_v1_8_deep_space_exploration(out_pptx)
    assert os.path.exists(deck_path)

    val = StaticValidator().validate(deck_path)
    assert val["is_valid"] is True, f"Defects: {val.get('defects')}"
    assert len(val["defects"]) == 0
    assert val["total_slides"] == 12

    # Check massive hierarchy: Slide 1 and Slide 10 have 100 stars
    import pptx
    prs_ds = pptx.Presentation(deck_path)
    assert any(len(s.shapes) >= 100 for s in prs_ds.slides)

    clean = verify_clean_powerpoint_open(deck_path)
    assert clean["is_clean"] is True, f"COM Clean open error: {clean.get('error')}"
