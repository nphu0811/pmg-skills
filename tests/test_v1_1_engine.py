"""
Automated tests for PMG Skills v1.1.0+:
- Engine-First Architecture & Backward Compatibility
- OpenXML Static Validation & Anti-Shortcut Enforcement
- ExpectedTimeline Extraction
- 3-Loop Self-Healing Engine
- Canonical CLI Subcommands
"""

import os
import json
import subprocess
import sys
import pytest
import pptx
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

from pmg import (
    VERSION,
    get_engine_provenance,
    ValidationStatus,
    MotionDeck,
    Themes,
    LightThemeCraftDeck,
    CinematicColors,
    ForgeDeck,
    Cinematic3DDeck,
    build_environmental_protection_deck,
    StaticValidator,
    extract_expected_timeline,
    execute_repair_loop,
    PMGOrchestrator,
)


def test_version_and_provenance():
    assert VERSION in ("1.1.0", "1.2.0", "1.3.0", "1.4.0", "1.5.0")
    prov = get_engine_provenance("test_suite")
    assert prov["engine_name"] == "pmg-skills"
    assert prov["engine_version"] == VERSION
    assert prov["canonical_entrypoint"] == "test_suite"
    assert "git_commit" in prov


def test_readme_and_skill_backward_compatibility(tmp_path):
    deck = LightThemeCraftDeck(title="Autonomous AI Core", theme=CinematicColors.LIGHT_ACADEMIC_CYAN)
    deck.add_craft_slide(
        title="NEURAL FABRIC ARCHITECTURE",
        category="AI COMPUTE PLATFORM",
        badge="STAGE 01",
        subtitle="Global overview of high-throughput heterogeneous compute clusters.",
    )
    deck.add_craft_slide(
        title="TENSOR ACCELERATOR ENGINE",
        category="DEEP LEARNING HARDWARE",
        badge="STAGE 02",
        subtitle="Sub-millisecond matrix execution running across distributed pipeline nodes.",
    )
    out_craft = str(tmp_path / "craft_test.pptx")
    deck.save(out_craft)
    assert os.path.exists(out_craft)

    val = StaticValidator().validate(out_craft)
    assert val["is_valid"] is True
    assert val["total_slides"] == 2

    c3d = Cinematic3DDeck(title="AI Neural Core")
    c3d.add_cinematic_slide(scene_idx=1, camera_zoom=1.0, camera_focus_center=(6.66, 3.75))
    c3d.add_cinematic_slide(scene_idx=2, camera_zoom=2.2, camera_focus_center=(4.2, 3.8))
    out_3d = str(tmp_path / "c3d_test.pptx")
    c3d.save(out_3d)
    assert os.path.exists(out_3d)
    val_3d = StaticValidator().validate(out_3d)
    assert val_3d["is_valid"] is True


def test_environmental_protection_10_slide_builder(tmp_path):
    out_pptx = str(tmp_path / "env_10_slides.pptx")
    build_environmental_protection_deck(out_pptx, num_slides=10)
    assert os.path.exists(out_pptx)

    validator = StaticValidator()
    res = validator.validate(out_pptx)
    assert res["is_valid"] is True
    assert res["total_slides"] == 10
    assert len(res["missing_transitions"]) == 0

    anchors = res["traveling_anchors"]
    assert "!!Traveling_Kite" in anchors
    assert anchors["!!Traveling_Kite"] == list(range(1, 11))
    assert "!!Earth_Globe" in anchors
    assert anchors["!!Earth_Globe"] == list(range(1, 11))
    assert "!!Meadow_Front" in anchors
    assert anchors["!!Meadow_Front"] == list(range(1, 11))

    timeline = extract_expected_timeline(out_pptx, hold_duration_ms=1200)
    assert timeline.total_slides == 10
    assert len(timeline.slides) == 10
    s2_ids = {o.object_id: o for o in timeline.slides[1].objects}
    assert "!!Traveling_Kite" in s2_ids
    assert s2_ids["!!Traveling_Kite"].trigger_type == "slide_transition"
    assert s2_ids["!!Traveling_Kite"].initial_position != s2_ids["!!Traveling_Kite"].final_position


def test_anti_shortcut_flattened_slide_detection(tmp_path):
    from PIL import Image
    img_path = str(tmp_path / "flat.png")
    Image.new("RGB", (400, 300), (100, 150, 200)).save(img_path)

    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_picture(img_path, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    flat_pptx = str(tmp_path / "flattened.pptx")
    prs.save(flat_pptx)

    res = StaticValidator().validate(flat_pptx)
    assert res["is_valid"] is False
    assert any("flattened slide shortcut is forbidden" in d["description"] for d in res["defects"])


def test_three_loop_self_healing_engine(tmp_path):
    deck = MotionDeck(theme=Themes.LIGHT_ECO_EMERALD)
    s1 = deck.add_slide(transition="none")
    deck.add_hero_title(s1, "SLIDE 1", subtitle="Start")
    deck.add_thematic_traveling_entity(s1, left=1.0, top=2.0, name="!!Traveling_Kite")

    s2 = deck.add_slide(transition="none")
    deck.add_hero_title(s2, "SLIDE 2", subtitle="Middle gap")
    deck.add_glass_card(s2, left=1.0, top=3.0, width=4.0, height=2.0, name="!!card")

    s3 = deck.add_slide(transition="none")
    deck.add_hero_title(s3, "SLIDE 3", subtitle="End")
    deck.add_thematic_traveling_entity(s3, left=9.0, top=2.0, name="!!Traveling_Kite")

    broken_pptx = str(tmp_path / "broken_deck.pptx")
    deck.save(broken_pptx)

    validator = StaticValidator()
    val0 = validator.validate(broken_pptx)
    assert val0["is_valid"] is False
    assert 2 in val0["missing_transitions"]
    assert "!!Traveling_Kite" in val0["discontinuous_anchors"]

    from pmg.models import ReviewDefect
    defects0 = [ReviewDefect(**d) for d in val0["defects"]]
    r1 = execute_repair_loop(broken_pptx, 1, defects0, val0)
    assert r1["loop_type"] == "targeted_repair"

    val1 = validator.validate(broken_pptx)
    assert len(val1["missing_transitions"]) == 0
    assert "!!Traveling_Kite" in val1["discontinuous_anchors"]

    defects1 = [ReviewDefect(**d) for d in val1["defects"]]
    r2 = execute_repair_loop(broken_pptx, 2, defects1, val1)
    assert r2["loop_type"] == "structural_repair"

    val2 = validator.validate(broken_pptx)
    assert val2["is_valid"] is True
    assert len(val2["discontinuous_anchors"]) == 0


def test_cli_commands(tmp_path):
    out_pptx = str(tmp_path / "cli_env.pptx")
    timeline_json = str(tmp_path / "timeline.json")
    recon_pptx = str(tmp_path / "reconstructed.pptx")

    res_build = subprocess.run(
        [sys.executable, "-m", "pmg", "build", "-o", out_pptx, "--preset", "environment", "--slides", "4", "--skip-playback"],
        capture_output=True,
        text=True,
    )
    assert res_build.returncode == 0
    assert os.path.exists(out_pptx)

    res_ins = subprocess.run(
        [sys.executable, "-m", "pmg", "inspect", out_pptx, "--timeline-out", timeline_json],
        capture_output=True,
        text=True,
    )
    assert res_ins.returncode == 0
    assert os.path.exists(timeline_json)

    res_val = subprocess.run(
        [sys.executable, "-m", "pmg", "validate", out_pptx, "--static-only"],
        capture_output=True,
        text=True,
    )
    assert res_val.returncode == 0

    res_rec = subprocess.run(
        [sys.executable, "-m", "pmg", "reconstruct", out_pptx, "-o", recon_pptx],
        capture_output=True,
        text=True,
    )
    assert res_rec.returncode == 0
    assert os.path.exists(recon_pptx)
