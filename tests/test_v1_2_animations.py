"""
Automated tests for PMG Skills v1.2.0:
- Declarative SlideAnimationTimeline Engine
- Native OpenXML <p:timing> Compilation & Round-Trip Extraction
- v1.2 Single-Slide Showcase Builder (>= 8 objects, >= 15 native animation effects)
- PowerPoint COM Animation Pane Verification (slide.TimeLine.MainSequence)
- Intra-Slide Timing Self-Healing (broken spTgt & invalid motion paths)
"""

import os
import pytest
import pptx
from pptx.util import Inches
from lxml import etree

from pmg import (
    SlideAnimationTimeline,
    parse_slide_timing_openxml,
    verify_animation_pane_com,
    build_v1_2_single_slide_animation_showcase,
    StaticValidator,
    extract_expected_timeline,
    execute_repair_loop,
    is_powerpoint_available,
    ReviewDefect,
)
from pmg.animations import P_NS, NSMAP


def test_declarative_timeline_scheduling_and_roundtrip(tmp_path):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    s1 = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(1))
    s1.name = "Alpha_Node"
    s2 = slide.shapes.add_shape(1, Inches(4), Inches(1), Inches(2), Inches(1))
    s2.name = "Beta_Node"

    tl = SlideAnimationTimeline(slide_number=1)
    e1 = tl.add_effect(s1, "appear", trigger="with_previous", duration_ms=200, delay_ms=50)
    e2 = tl.add_effect(s2, "fade", trigger="with_previous", duration_ms=500, delay_ms=100)
    e3 = tl.add_effect(s1, "motion_path", trigger="after_previous", duration_ms=800, motion_path="M 0 0 C 0.1 -0.1 0.2 0.1 0.3 0 E")
    e4 = tl.add_effect(s2, "rotate", trigger="on_click", duration_ms=600, rotation_deg=360.0)
    e5 = tl.add_effect(s1, "scale", trigger="with_previous", duration_ms=400, scale_factor=(1.25, 1.25), auto_reverse=True, repeat_count=2.0)

    # Verify deterministic schedule calculation
    assert e1.computed_start_ms == 50
    assert e1.computed_end_ms == 250
    assert e2.computed_start_ms == 100
    assert e2.computed_end_ms == 600
    assert e3.computed_start_ms == 600
    assert e3.computed_end_ms == 1400
    assert e4.computed_start_ms == 1650  # 1400 + 250ms click gap
    assert e4.computed_end_ms == 2250
    assert e5.computed_start_ms == 1650
    assert e5.computed_end_ms == 1650 + (400 * 2 * 2)  # repeat=2, auto_reverse=True -> 1600ms

    tl.compile_to_openxml(slide)
    out_pptx = str(tmp_path / "tl_roundtrip.pptx")
    prs.save(out_pptx)

    # Parse back via parse_slide_timing_openxml
    prs2 = pptx.Presentation(out_pptx)
    specs, clicks = parse_slide_timing_openxml(prs2.slides[0], slide_number=1)
    assert len(specs) == 5
    assert len(clicks) == 1
    assert clicks[0]["target_object"] == "Beta_Node"


def test_v1_2_single_slide_showcase_requirements(tmp_path):
    out_pptx = str(tmp_path / "v1_2_showcase.pptx")
    build_v1_2_single_slide_animation_showcase(out_pptx)
    assert os.path.exists(out_pptx)

    validator = StaticValidator()
    val = validator.validate(out_pptx)
    assert val["is_valid"] is True
    assert val["total_slides"] == 1

    s_info = val["slides"][0]
    assert s_info["editable_object_count"] >= 8
    assert s_info["native_animation_count"] >= 15

    # Verify ExpectedTimeline extraction includes all intra-slide effects
    timeline = extract_expected_timeline(out_pptx, hold_duration_ms=1200)
    assert timeline.total_slides == 1
    s1_spec = timeline.slides[0]
    intra_objs = [o for o in s1_spec.objects if o.trigger_type not in ("static", "slide_transition")]
    assert len(intra_objs) >= 15

    # Verify presence of parallel (with_previous), sequential (after_previous / on_click),
    # curved motion path, rotation, scale, entrance, and exit effects
    triggers = {o.trigger_type for o in intra_objs}
    assert "with_previous" in triggers
    assert "after_previous" in triggers
    assert "on_click" in triggers

    has_curved_path = any(o.motion_path and "C" in o.motion_path for o in intra_objs)
    has_rotation = any(abs(o.rotation_params.get("end_deg", 0.0)) >= 180.0 for o in intra_objs)
    has_scale = any(o.initial_size != o.final_size for o in intra_objs)
    has_entrance = any(o.opacity_params.get("start_alpha") == 0.0 and o.opacity_params.get("end_alpha") == 1.0 for o in intra_objs)
    has_exit = any(o.opacity_params.get("start_alpha") == 1.0 and o.opacity_params.get("end_alpha") == 0.0 for o in intra_objs)

    assert has_curved_path is True
    assert has_rotation is True
    assert has_scale is True
    assert has_entrance is True
    assert has_exit is True

    # Verify in Microsoft PowerPoint COM Animation Pane if available
    if is_powerpoint_available():
        com_res = verify_animation_pane_com(out_pptx)
        assert com_res["verified_via_com"] is True
        assert com_res["total_main_sequence_effects"] >= 15


def test_intra_slide_timing_self_healing(tmp_path):
    out_pptx = str(tmp_path / "broken_timing.pptx")
    build_v1_2_single_slide_animation_showcase(out_pptx)

    # Deliberately corrupt a <p:spTgt> spid and a <p:animMotion> path
    prs = pptx.Presentation(out_pptx)
    timing = prs.slides[0]._element.find(f"{{{P_NS}}}timing")
    sp_tgts = timing.xpath(".//p:spTgt")
    sp_tgts[0].set("spid", "99999")
    anim_motions = timing.xpath(".//p:animMotion")
    anim_motions[0].set("path", "INVALID_PATH")
    prs.save(out_pptx)

    validator = StaticValidator()
    val0 = validator.validate(out_pptx)
    assert val0["is_valid"] is False
    hints = {d["remediation_hint"] for d in val0["defects"]}
    assert "repair_animation_target" in hints
    assert "repair_motion_path" in hints

    defects0 = [ReviewDefect(**d) for d in val0["defects"]]
    rep = execute_repair_loop(out_pptx, 1, defects0, val0)
    assert len(rep["actions_taken"]) >= 2

    val1 = validator.validate(out_pptx)
    assert val1["is_valid"] is True
