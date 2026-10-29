"""
tests/test_v2_1_0_universal_runtime.py
======================================
Automated Acceptance, Regression, and Integration Tests for PMG Skills v2.1.0:
Universal Runtime & Cross-Platform Reliability.

Covers:
- Universal capability detection (Windows w/ w/o PowerPoint, Linux, macOS, offline sandbox)
- Offline-first 6-tier dependency resolution strategy
- Safe execution without optional dependencies (cv2, scipy, pywin32)
- Consolidated Native OOXML Compiler (without requiring Microsoft PowerPoint)
- Comprehensive native PPTX fixture (editable shapes, morph, sequential & parallel effects, motion paths)
- Three-Level Validation:
  * Level 1: Structural (ZIP integrity, XML parsing, OPC types, relationships, shape IDs, target existence)
  * Level 2: Visual (rendered preview reporting vs SKIPPED when renderer is absent)
  * Level 3: Playback (live PowerPoint execution vs NOT_TESTED when PowerPoint is absent)
- Repair prevention & negative fixtures (broken targets, duplicate IDs, missing content types)
- Unified Public API (`PMGRuntime.auto()`, `runtime.build()`, `runtime.validate()`, `runtime.doctor()`)
- CLI commands (`pmg doctor`, `pmg capabilities`, `pmg build`, `pmg validate`)
- Remote worker Mode C adapter
- Real PowerPoint Desktop integration test when available on host
"""

import os
import sys
import json
import zipfile
import subprocess
import xml.etree.ElementTree as ET
from unittest.mock import patch, MagicMock
import pytest
import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from pmg.version import VERSION
from pmg.models import ValidationStatus, ReviewDefect
from pmg.capabilities import (
    CapabilityState,
    ExecutionMode,
    EnvironmentCapabilities,
    detect_environment,
)
from pmg.dependencies import (
    resolve_dependencies,
    find_compatible_offline_wheels,
)
from pmg.compiler import (
    NativePPTXCompiler,
    CompilerReport,
)
from pmg.validation import (
    validate_level1_structural,
    validate_level2_visual,
    validate_level3_playback,
    validate_presentation,
    ComprehensiveValidationReport,
    Level1Result,
    Level2Result,
    Level3Result,
)
from pmg.runtime import (
    PMGRuntime,
    PMGBuildResult,
)
from pmg.validator import StaticValidator, NS
from pmg.worker.client import PMGWorkerClient


# ==========================================
# 1. VERSION & PROVENANCE
# ==========================================

def test_v2_1_0_version():
    assert VERSION in ("2.1.0", "2.2.0", "2.3.0", "2.4.0")


# ==========================================
# 2. UNIVERSAL CAPABILITY DETECTION
# ==========================================

def test_detect_environment_host():
    """Verifies deterministic host detection in the current environment."""
    caps = detect_environment()
    assert isinstance(caps, EnvironmentCapabilities)
    assert caps.python_available is True
    assert caps.can_compile_native_pptx is True
    assert "native_compiler" in caps.available_backends
    assert isinstance(caps.format_diagnostic_report(), str)
    assert "PMG Skills Universal Runtime" in caps.format_diagnostic_report()


def test_detect_environment_mock_linux_sandbox():
    """Simulates a restricted Linux AI web sandbox (no PowerPoint, no net, pure python)."""
    mock_probes = {
        "platform": "linux",
        "os_name": "Linux",
        "python_available": True,
        "pptx_available": True,
        "lxml_available": True,
        "pillow_available": True,
        "powerpoint_available": False,
        "com_available": False,
        "renderer_available": False,
        "network_available": False,
        "remote_worker_available": False,
        "cv2_available": False,
        "scipy_available": False,
    }
    caps = detect_environment(probes_override=mock_probes)
    assert caps.platform == "linux"
    assert caps.execution_mode == ExecutionMode.WEB_SANDBOX
    assert caps.powerpoint_available is False
    assert caps.can_compile_native_pptx is True
    assert caps.can_validate_structural is True
    assert caps.can_validate_visual is False
    assert caps.can_validate_playback is False
    assert caps.available_backends == ["native_compiler"]


def test_detect_environment_mock_macos():
    """Simulates macOS local desktop environment."""
    mock_probes = {
        "platform": "darwin",
        "os_name": "Darwin",
        "powerpoint_available": False,
        "com_available": False,
        "renderer_available": False,
        "network_available": True,
    }
    caps = detect_environment(probes_override=mock_probes)
    assert caps.platform == "darwin"
    assert caps.execution_mode == ExecutionMode.WEB_SANDBOX
    assert caps.com_available is False
    assert caps.can_compile_native_pptx is True


def test_detect_environment_mock_windows_no_powerpoint():
    """Simulates a Windows workstation without Microsoft Office installed."""
    mock_probes = {
        "platform": "win32",
        "os_name": "Windows",
        "powerpoint_available": False,
        "com_available": True,
        "renderer_available": False,
        "network_available": True,
    }
    caps = detect_environment(probes_override=mock_probes)
    assert caps.platform == "win32"
    assert caps.powerpoint_available is False
    assert caps.can_compile_native_pptx is True
    assert caps.can_validate_playback is False


def test_detect_environment_mock_remote_worker_configured():
    """Simulates Mode C when a remote Windows validation worker is configured & reachable."""
    mock_probes = {
        "platform": "linux",
        "os_name": "Linux",
        "remote_worker_url": "https://remote-worker.internal:8765",
        "remote_worker_available": True,
        "powerpoint_available": False,
    }
    caps = detect_environment(probes_override=mock_probes)
    assert caps.remote_worker_configured is True
    assert caps.remote_worker_available is True
    assert caps.execution_mode == ExecutionMode.REMOTE_WORKER
    assert "remote_windows_worker" in caps.available_backends
    assert caps.can_validate_playback is True


# ==========================================
# 3. OFFLINE-FIRST DEPENDENCY STRATEGY
# ==========================================

def test_offline_dependency_resolution_strategy(tmp_path):
    """Verifies the 6-tier dependency resolution sequence and cached wheel discovery."""
    # Create mock cached wheel directory
    cache_dir = tmp_path / "offline_cache"
    cache_dir.mkdir()
    mock_wheel = cache_dir / "python_pptx-0.6.21-py3-none-any.whl"
    mock_wheel.write_bytes(b"PK\x03\x04mockwheel")

    # Find wheels
    found = find_compatible_offline_wheels(str(cache_dir))
    assert len(found) == 1
    assert "python_pptx" in found[0]

    # Resolve dependencies
    res = resolve_dependencies(allow_network=False, cache_override_dir=str(cache_dir))
    assert "python-pptx" in res
    assert res["python-pptx"].status == "installed"
    assert res["python-pptx"].compatible is True


# ==========================================
# 4. SHARED NATIVE OOXML COMPILER (WITHOUT POWERPOINT)
# ==========================================

def test_compiler_native_pptx_from_scratch(tmp_path):
    """
    Compiles a native presentation from scratch using programmatic API:
    - 16:9 canvas geometry
    - Light theme
    - Frosted glassmorphism card
    - Native text box
    - Morph transition
    - Intra-slide Animation Pane timing tree (entrance fade)
    Verifies that NO PowerPoint Desktop is launched or required.
    """
    output_pptx = str(tmp_path / "compiled_programmatic.pptx")
    compiler = NativePPTXCompiler()
    compiler.new_presentation()

    # Slide 1
    s1 = compiler.add_slide(transition="fade", duration_ms=1000)
    card1 = compiler.add_glass_card(
        slide=s1,
        left_in=1.0,
        top_in=2.0,
        width_in=5.0,
        height_in=3.0,
        name="!!HeroCard",
        title="Frosted Card",
        body="Compiled natively without Microsoft PowerPoint.",
    )
    compiler.add_animation(
        slide=s1,
        shape=card1,
        effect_type="fade",
        trigger="with_previous",
        duration_ms=800,
        delay_ms=200,
    )

    # Slide 2 with Morph
    s2 = compiler.add_slide(transition="morph", duration_ms=1400)
    compiler.add_glass_card(
        slide=s2,
        left_in=6.0,
        top_in=2.0,
        width_in=6.0,
        height_in=3.5,
        name="!!HeroCard",
        title="Morphed Hero Card",
        body="Persistent object identity for 60 FPS Morph interpolation.",
    )

    rep = compiler.save(output_pptx)
    assert os.path.exists(output_pptx)
    assert rep.total_slides == 2
    assert rep.total_shapes >= 2
    assert rep.total_animations == 1
    assert rep.morph_transition_count >= 1
    assert rep.is_fully_native is True
    assert "!!HeroCard" in rep.named_anchors


def test_compiler_from_json_specification(tmp_path):
    """Verifies compilation from a structured JSON presentation specification."""
    spec = {
        "theme": "LIGHT_ACADEMIC_CYAN",
        "slides": [
            {
                "title": "Quantum Entanglement",
                "subtitle": "Cross-Slide 60 FPS Morph Continuity",
                "transition": "fade",
                "duration_ms": 1200,
                "elements": [
                    {
                        "type": "card",
                        "name": "!!Qubit_Node",
                        "left": 1.0,
                        "top": 2.5,
                        "width": 5.0,
                        "height": 3.0,
                        "title": "Superconducting Node",
                        "body": "Phase coherence maintained across multi-core processors.",
                    }
                ],
                "animations": [
                    {
                        "target": "!!Qubit_Node",
                        "effect": "fade",
                        "trigger": "with_previous",
                        "duration_ms": 750,
                    }
                ],
            },
            {
                "title": "Scaled Qubit Topology",
                "transition": "morph",
                "duration_ms": 1400,
                "elements": [
                    {
                        "type": "card",
                        "name": "!!Qubit_Node",
                        "left": 4.0,
                        "top": 2.0,
                        "width": 8.0,
                        "height": 4.0,
                        "title": "Scaled Topology Array",
                    }
                ],
            },
        ],
    }
    output_pptx = str(tmp_path / "compiled_from_spec.pptx")
    compiler = NativePPTXCompiler()
    rep = compiler.compile_from_spec(spec, output_pptx)

    assert os.path.exists(output_pptx)
    assert rep.total_slides == 2
    assert rep.total_animations == 1

    # Verify OpenXML content without PowerPoint
    with zipfile.ZipFile(output_pptx, "r") as zf:
        s1_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        s2_xml = zf.read("ppt/slides/slide2.xml").decode("utf-8")

    assert "p:timing" in s1_xml
    assert "fade" in s1_xml
    assert "p159:morph" in s2_xml


def test_native_pptx_fixture_with_all_motion_constructs(tmp_path):
    """
    Creates and validates a rich native presentation fixture containing:
    - Multiple editable objects (cards, text, geometric shapes)
    - 60 FPS Morph transitions across slides
    - Sequential entrance effects (`after_previous`)
    - Parallel animation effects (`with_previous`)
    - Motion paths (`motion_path` with cubic Bezier)
    - Exit effects (`fade_out`)
    """
    output_pptx = str(tmp_path / "rich_motion_fixture.pptx")
    compiler = NativePPTXCompiler()
    compiler.new_presentation()

    # Slide 1: Entrances, Parallel Effects, and Motion Path
    s1 = compiler.add_slide(transition="fade", duration_ms=1000)
    title = compiler.add_textbox(s1, 0.8, 0.8, 8.0, 1.0, "KINEMATICS & ORBITAL DYNAMICS", name="!!Title", font_size_pt=32, font_bold=True)
    hero = compiler.add_shape(s1, MSO_SHAPE.OVAL, 1.0, 2.5, 2.0, 2.0, name="!!Satellite", text="SAT-1")
    card_a = compiler.add_glass_card(s1, 4.0, 2.5, 4.0, 2.0, name="!!TelemetryA", title="Velocity", body="7.8 km/s")
    card_b = compiler.add_glass_card(s1, 8.5, 2.5, 4.0, 2.0, name="!!TelemetryB", title="Altitude", body="408 km")

    # Animation 1: Title Fade In
    compiler.add_animation(s1, title, effect_type="fade", trigger="with_previous", duration_ms=600)
    # Animation 2: Hero Satellite Fly In
    compiler.add_animation(s1, hero, effect_type="fly_in", direction="up", trigger="after_previous", duration_ms=700)
    # Animation 3: Satellite Motion Path (parallel with Telemetry A)
    compiler.add_animation(s1, hero, effect_type="motion_path", motion_path="M 0 0 C 0.1 0.05 0.2 -0.05 0.3 0.1 E", trigger="after_previous", duration_ms=900)
    # Animation 4: Telemetry A Fade In (parallel)
    compiler.add_animation(s1, card_a, effect_type="fade", trigger="with_previous", duration_ms=600, delay_ms=100)
    # Animation 5: Telemetry B Zoom In (sequential after Telemetry A)
    compiler.add_animation(s1, card_b, effect_type="zoom", trigger="after_previous", duration_ms=600)

    # Slide 2: Morph Transition + Exit
    s2 = compiler.add_slide(transition="morph", duration_ms=1600)
    compiler.add_textbox(s2, 0.8, 0.8, 8.0, 1.0, "ORBITAL INSERTION COMPLETE", name="!!Title", font_size_pt=32, font_bold=True)
    hero_s2 = compiler.add_shape(s2, MSO_SHAPE.OVAL, 6.0, 2.0, 3.5, 3.5, name="!!Satellite", text="SAT-1 (GEO)")
    compiler.add_glass_card(s2, 1.0, 2.5, 4.5, 3.5, name="!!TelemetryA", title="Final Orbit", body="Geostationary Transfer")

    # Animation on Slide 2: Satellite Fade Out
    compiler.add_animation(s2, hero_s2, effect_type="fade_out", trigger="after_previous", duration_ms=800, delay_ms=500)

    rep = compiler.save(output_pptx)
    assert rep.total_slides == 2
    assert rep.total_animations == 6
    assert rep.morph_transition_count >= 1

    # Structural inspection
    l1 = validate_level1_structural(output_pptx)
    assert l1.status == ValidationStatus.PASS
    assert l1.zip_integrity is True
    assert l1.opc_content_types_valid is True
    assert l1.relationships_valid is True
    assert l1.unique_identifiers_valid is True
    assert l1.animation_targets_valid is True
    assert l1.timing_graph_valid is True
    assert l1.morph_metadata_consistent is True


# ==========================================
# 5. THREE-LEVEL VALIDATION & HONEST REPORTING
# ==========================================

def test_three_level_validation_sandbox_reporting(tmp_path):
    """
    Verifies that in a sandbox environment lacking PowerPoint:
    - Level 1 reports PASS
    - Level 2 reports SKIPPED — No slide rendering backend available
    - Level 3 reports NOT_TESTED — PowerPoint Desktop unavailable
    - NEVER reports Level 3 PASS without actual PowerPoint Desktop!
    """
    pptx_path = str(tmp_path / "sandbox_test.pptx")
    compiler = NativePPTXCompiler()
    compiler.new_presentation()
    s = compiler.add_slide(transition="fade")
    compiler.add_textbox(s, 1.0, 1.0, 5.0, 1.0, "Sandbox Test")
    compiler.save(pptx_path)

    # Force sandbox capabilities (no PowerPoint, no renderer)
    sandbox_caps = detect_environment(probes_override={
        "powerpoint_available": False,
        "com_available": False,
        "renderer_available": False,
    })

    rep = validate_presentation(pptx_path, capabilities=sandbox_caps)
    assert rep.level1.status == ValidationStatus.PASS
    assert rep.level2.status == ValidationStatus.SKIPPED
    assert "SKIPPED" in rep.level2.summary
    assert rep.level3.status == ValidationStatus.NOT_TESTED
    assert "NOT_TESTED" in rep.level3.summary
    assert rep.overall_status == ValidationStatus.PASS
    assert rep.can_deliver is True
    assert "NOT_TESTED" in rep.summary


def test_level1_validation_catches_broken_animation_target(tmp_path):
    """
    Regression test: ensures Level 1 detects dangling `<p:spTgt>` references that
    target non-existent shape IDs (which trigger PowerPoint Repair warnings).
    """
    bad_pptx = str(tmp_path / "broken_target.pptx")
    compiler = NativePPTXCompiler()
    compiler.new_presentation()
    s = compiler.add_slide()
    compiler.add_textbox(s, 1.0, 1.0, 4.0, 1.0, "Test Box")
    compiler.save(bad_pptx)

    # Inject dangling animation target spid="99999" into slide1.xml
    with zipfile.ZipFile(bad_pptx, "r") as zf:
        namelist = zf.namelist()
        parts = {p: zf.read(p) for p in namelist}

    s1_xml = parts["ppt/slides/slide1.xml"].decode("utf-8")
    dangling_timing = (
        f'<p:timing xmlns:p="{NS["p"]}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        f'<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst><p:par><p:cTn id="4" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="withEffect">'
        f'<p:childTnLst><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="5" dur="800"/>'
        f'<p:tgtEl><p:spTgt spid="99999"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst>'
        f'</p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:seq>'
        f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )
    # Insert before closing </p:sld>
    modified_xml = s1_xml.replace("</p:sld>", f"{dangling_timing}</p:sld>")
    parts["ppt/slides/slide1.xml"] = modified_xml.encode("utf-8")

    with zipfile.ZipFile(bad_pptx, "w") as zf:
        for name, data in parts.items():
            zf.writestr(name, data)

    # Verify Level 1 catches the defect
    l1 = validate_level1_structural(bad_pptx)
    assert l1.status == ValidationStatus.FAIL
    assert l1.animation_targets_valid is False
    assert any("99999" in d.description for d in l1.defects)


def test_level1_validation_catches_duplicate_shape_ids(tmp_path):
    """
    Regression test: ensures Level 1 detects duplicate shape IDs on a slide.
    """
    dup_pptx = str(tmp_path / "duplicate_id.pptx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t1 = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    t1.text = "Box 1"
    t2 = s.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    t2.text = "Box 2"
    prs.save(dup_pptx)

    # Force duplicate shape id="2"
    with zipfile.ZipFile(dup_pptx, "r") as zf:
        namelist = zf.namelist()
        parts = {p: zf.read(p) for p in namelist}

    s1_str = parts["ppt/slides/slide1.xml"].decode("utf-8")
    # Replace second shape id="3" with id="2"
    s1_dup = s1_str.replace('id="3"', 'id="2"')
    parts["ppt/slides/slide1.xml"] = s1_dup.encode("utf-8")

    with zipfile.ZipFile(dup_pptx, "w") as zf:
        for name, data in parts.items():
            zf.writestr(name, data)

    l1 = validate_level1_structural(dup_pptx)
    assert l1.status == ValidationStatus.FAIL
    assert l1.unique_identifiers_valid is False
    assert any("duplicate" in d.description.lower() for d in l1.defects)


def test_level1_validation_catches_missing_content_types(tmp_path):
    """
    Regression test: ensures Level 1 flags package parts lacking a declared ContentType.
    """
    bad_opc = str(tmp_path / "missing_ct.pptx")
    prs = pptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(bad_opc)

    # Inject undeclared part into ZIP without updating [Content_Types].xml
    with zipfile.ZipFile(bad_opc, "a") as zf:
        zf.writestr("ppt/undeclared_data.customxyz", b"DATA")

    l1 = validate_level1_structural(bad_opc)
    assert l1.status == ValidationStatus.FAIL
    assert any("has no registered ContentType" in d.description for d in l1.defects)


# ==========================================
# 6. UNIFIED PUBLIC API (`PMGRuntime`)
# ==========================================

def test_unified_api_pmg_runtime_build_and_validate(tmp_path):
    """
    Verifies target API usage:
    ```python
    from pmg import PMGRuntime
    runtime = PMGRuntime.auto()
    result = runtime.build(project=..., output=..., validate=True)
    ```
    """
    out_pptx = str(tmp_path / "runtime_deck.pptx")
    runtime = PMGRuntime.auto()

    spec = {
        "title": "Quantum Photonics",
        "theme": "LIGHT_ACADEMIC_CYAN",
        "slides": [
            {
                "title": "Waveguide Routing",
                "elements": [
                    {
                        "type": "card",
                        "name": "!!Waveguide",
                        "left": 1.0,
                        "top": 2.5,
                        "width": 5.0,
                        "height": 3.0,
                        "title": "Silicon Nitride Array",
                    }
                ],
                "animations": [
                    {"target": "!!Waveguide", "effect": "fade", "duration_ms": 700}
                ],
            }
        ],
    }

    res = runtime.build(project=spec, output=out_pptx, validate=True)
    assert isinstance(res, PMGBuildResult)
    assert os.path.exists(res.output_path)
    assert res.is_valid is True
    assert res.can_deliver is True
    assert res.validation_report is not None
    assert res.validation_report.level1.status == ValidationStatus.PASS
    # Verify companion manifest was generated
    companion_json = f"{out_pptx}.validation_report.json"
    assert os.path.exists(companion_json)


def test_unified_api_validate_standalone(tmp_path):
    """Verifies runtime.validate(pptx_path)."""
    deck_path = str(tmp_path / "standalone.pptx")
    compiler = NativePPTXCompiler()
    compiler.new_presentation()
    compiler.add_slide()
    compiler.save(deck_path)

    runtime = PMGRuntime.auto()
    report = runtime.validate(deck_path)
    assert isinstance(report, ComprehensiveValidationReport)
    assert report.level1.status == ValidationStatus.PASS


# ==========================================
# 7. CLI COMMANDS (`doctor`, `capabilities`, `build`, `validate`)
# ==========================================

def test_cli_doctor():
    """Runs `pmg doctor` and confirms exit code 0."""
    res = subprocess.run([sys.executable, "-m", "pmg", "doctor"], capture_output=True, text=True, timeout=15)
    assert res.returncode == 0
    assert "PMG Skills Universal Runtime" in res.stdout
    assert "Native Compilation: READY" in res.stdout


def test_cli_capabilities():
    """Runs `pmg capabilities` and verifies valid JSON output."""
    res = subprocess.run([sys.executable, "-m", "pmg", "capabilities"], capture_output=True, text=True, timeout=15)
    assert res.returncode == 0
    data = json.loads(res.stdout)
    assert "platform" in data
    assert "can_compile_native_pptx" in data
    assert data["can_compile_native_pptx"] is True


def test_cli_build_and_validate_from_spec(tmp_path):
    """Runs `pmg build spec.json -o out.pptx` via command-line."""
    spec_file = str(tmp_path / "cli_spec.json")
    out_pptx = str(tmp_path / "cli_out.pptx")

    spec_data = {
        "slides": [
            {
                "title": "CLI Integration Test",
                "elements": [{"type": "card", "name": "!!TestCard", "title": "CLI Card"}],
                "animations": [{"target": "!!TestCard", "effect": "fade"}],
            }
        ]
    }
    with open(spec_file, "w", encoding="utf-8") as f:
        json.dump(spec_data, f)

    cmd = [sys.executable, "-m", "pmg", "build", spec_file, "-o", out_pptx, "--skip-playback"]
    p_build = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
    assert p_build.returncode == 0
    assert os.path.exists(out_pptx)

    # Run `pmg validate` on the output
    p_val = subprocess.run([sys.executable, "-m", "pmg", "validate", out_pptx], capture_output=True, text=True, timeout=30)
    assert p_val.returncode == 0
    assert "PMG Three-Level Validation" in p_val.stdout


# ==========================================
# 8. REMOTE WORKER ADAPTER (MODE C)
# ==========================================

def test_remote_worker_compile_adapter_unconfigured(tmp_path):
    """Verifies that an unconfigured remote worker fails gracefully to local execution."""
    client = PMGWorkerClient(worker_url="", worker_token=None)
    out_file = str(tmp_path / "remote_unconfigured.pptx")
    ok = client.compile_remote({"slides": []}, out_file)
    assert ok is False
    assert not os.path.exists(out_file)


# ==========================================
# 9. REAL POWERPOINT DESKTOP INTEGRATION (WHEN AVAILABLE)
# ==========================================

def test_real_powerpoint_desktop_clean_open_if_available(tmp_path):
    """
    If Microsoft PowerPoint Desktop is installed on this machine,
    verifies clean COM opening (`DisplayAlerts = ppAlertsAll`, `OpenAndRepair = False`)
    with zero repair warnings.
    """
    caps = detect_environment()
    if not (caps.platform == "win32" and caps.powerpoint_available and caps.com_available):
        pytest.skip("PowerPoint Desktop COM is not installed in this environment.")

    test_deck = str(tmp_path / "real_ppt_open.pptx")
    compiler = NativePPTXCompiler()
    compiler.new_presentation()
    s = compiler.add_slide(transition="morph", duration_ms=1200)
    card = compiler.add_glass_card(s, 1.0, 2.0, 5.0, 3.0, name="!!Hero", title="Real Desktop Test")
    compiler.add_animation(s, card, effect_type="fade", trigger="with_previous", duration_ms=800)
    compiler.save(test_deck)

    from pmg.worker.powerpoint_probe import verify_clean_powerpoint_open
    clean_res = verify_clean_powerpoint_open(test_deck)
    assert clean_res["is_clean"] is True
    assert clean_res["display_alerts"] == 2
    assert len(clean_res["interruptions"]) == 0
